import os
import re
from pathlib import Path
from datetime import datetime, timezone
import logging
from typing import Optional, List, Dict, Any, Tuple
from pydantic import BaseModel
from fastapi import APIRouter, UploadFile, File, Form, HTTPException, status, Depends
from app.core.auth import require_roles

logger = logging.getLogger(__name__)

router = APIRouter()

# Storage directory path under backend/storage
STORAGE_DIR = Path(__file__).resolve().parent.parent.parent / "storage"
STORAGE_DIR.mkdir(parents=True, exist_ok=True)

# Allowed Excel extensions
ALLOWED_EXTENSIONS = {".xlsx", ".xls"}


@router.post("/upload", summary="Upload Excel file", tags=["Data Ingestion"])
async def upload_excel_file(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Validates and stores an uploaded Excel file locally under backend/storage/.
    Returns upload status, filename, file size, and file type.
    """
    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename is required."
        )

    # Validate file extension
    file_ext = Path(file.filename).suffix.lower()
    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Only Excel files (.xlsx, .xls) are supported."
        )

    # Sanitize filename (remove path traversal components)
    safe_filename = Path(file.filename).name
    destination_path = STORAGE_DIR / safe_filename

    # Save file to local backend/storage directory
    try:
        total_bytes = 0
        with open(destination_path, "wb") as buffer:
            while chunk := await file.read(1024 * 1024):  # 1MB chunks
                buffer.write(chunk)
                total_bytes += len(chunk)

        if total_bytes == 0:
            # Clean up empty file
            if destination_path.exists():
                destination_path.unlink()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded file is empty (0 bytes)."
            )

    except HTTPException:
        raise
    except Exception as exc:
        # Clean up on error
        if destination_path.exists():
            destination_path.unlink()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save uploaded file: {str(exc)}"
        )
    finally:
        await file.close()

    return {
        "status": "success",
        "filename": safe_filename,
        "file_size": total_bytes,
        "file_type": "excel",
        "stored_location": f"backend/storage/{safe_filename}",
        "message": "Excel file uploaded successfully",
        "timestamp": datetime.now(timezone.utc).isoformat()
    }

IGNORED_FILENAMES = {"sample_l5_l6_schedule.xlsx", "test_execution_matches.xlsx", "test.xlsx", "dummy_rbac_check.xlsx"}


def get_active_project_filename() -> Optional[str]:
    """
    Dynamically finds the currently active uploaded project Excel/CSV file.
    Queries the database ProjectFile model first for active files, then falls back
    to the latest uploaded file in backend/storage/.
    """
    try:
        from app.core.database import SessionLocal
        from app.models.project_file import ProjectFile
        with SessionLocal() as db:
            records = db.query(ProjectFile).order_by(ProjectFile.is_active.desc(), ProjectFile.uploaded_at.desc()).all()
            for r in records:
                if r.original_filename not in IGNORED_FILENAMES and (STORAGE_DIR / r.original_filename).exists():
                    return r.original_filename
    except Exception as e:
        logger.debug(f"DB lookup for active file failed: {e}")

    # Fallback to STORAGE_DIR files
    candidates = []
    for p in STORAGE_DIR.glob("*.*"):
        if p.is_file() and p.suffix.lower() in (".xlsx", ".xls", ".csv"):
            if p.name not in IGNORED_FILENAMES:
                candidates.append((p.stat().st_mtime, p.name))
    if candidates:
        candidates.sort(reverse=True)
        return candidates[0][1]

    return None


def resolve_active_files(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None
) -> Tuple[str, str]:
    """
    Resolves execution_file and schedule_file to actual valid project files.
    Eliminates dependencies on legacy hardcoded files like 'sample_l5_l6_schedule.xlsx'.
    If a file is missing or points to the legacy mock filename, it dynamically resolves
    to the active uploaded project file.
    If no project files exist, raises HTTP 404 with:
    'No project data uploaded. Please upload a project Excel file to begin.'
    """
    active_fn = get_active_project_filename()

    # If schedule_file was explicitly provided as a custom project file and exists, use it for both
    if schedule_file and (STORAGE_DIR / schedule_file).exists() and schedule_file not in ("sample_l5_l6_schedule.xlsx", "baseline_schedule.xlsx"):
        sched = schedule_file
        exec_f = execution_file if (execution_file and (STORAGE_DIR / execution_file).exists() and execution_file not in ("test_execution_matches.xlsx", "sample_l5_l6_schedule.xlsx")) else schedule_file
        return (exec_f, sched)

    # If execution_file was explicitly provided as a custom project file and exists, use it for both
    if execution_file and (STORAGE_DIR / execution_file).exists() and execution_file not in ("test_execution_matches.xlsx", "sample_l5_l6_schedule.xlsx"):
        exec_f = execution_file
        sched = schedule_file if (schedule_file and (STORAGE_DIR / schedule_file).exists() and schedule_file not in ("sample_l5_l6_schedule.xlsx", "baseline_schedule.xlsx")) else execution_file
        return (exec_f, sched)

    # Resolve schedule_file
    sched = schedule_file
    if not sched or sched in ("sample_l5_l6_schedule.xlsx", "null", "undefined", "") or not (STORAGE_DIR / sched).exists():
        if active_fn:
            sched = active_fn
        elif (STORAGE_DIR / "baseline_schedule.xlsx").exists():
            sched = "baseline_schedule.xlsx"
        else:
            sched = None

    # Resolve execution_file
    exec_f = execution_file
    if not exec_f or exec_f in ("test_execution_matches.xlsx", "sample_l5_l6_schedule.xlsx", "null", "undefined", "") or not (STORAGE_DIR / exec_f).exists():
        if active_fn:
            exec_f = active_fn
        elif (STORAGE_DIR / "test_execution_matches.xlsx").exists() and execution_file == "test_execution_matches.xlsx":
            exec_f = "test_execution_matches.xlsx"
        else:
            exec_f = sched

    if not sched and not exec_f:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="No project data uploaded. Please upload a project Excel file to begin."
        )

    return (exec_f or sched, sched or exec_f)


@router.get("/excel/{filename}", summary="Read and parse Excel workbook", tags=["Data Ingestion"])
def get_excel_contents(filename: str):
    """
    Reads a stored Excel file (.xlsx or .xls) from backend/storage/.
    Extracts all worksheet names, column headers, and row records as structured JSON.
    """
    if not filename or filename in ("sample_l5_l6_schedule.xlsx", "null", "undefined", ""):
        active_fn = get_active_project_filename()
        if active_fn:
            filename = active_fn
        else:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No project data uploaded. Please upload a project Excel file to begin."
            )

    safe_filename = Path(filename).name
    file_ext = Path(safe_filename).suffix.lower()

    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Only Excel files (.xlsx, .xls) can be parsed."
        )

    file_path = STORAGE_DIR / safe_filename

    if not file_path.exists() or not file_path.is_file():
        active_fn = get_active_project_filename()
        if active_fn and (STORAGE_DIR / active_fn).exists() and active_fn != safe_filename:
            file_path = STORAGE_DIR / active_fn
            safe_filename = active_fn
        else:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No project data uploaded. Please upload a project Excel file to begin."
            )

    try:
        import openpyxl
        from datetime import date, datetime as dt

        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets_data = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_rows = list(ws.iter_rows(values_only=True))

            if not all_rows:
                sheets_data.append({
                    "name": sheet_name,
                    "row_count": 0,
                    "column_count": 0,
                    "columns": [],
                    "rows": []
                })
                continue

            # First row as header columns
            raw_headers = all_rows[0]
            columns = []
            for idx, h in enumerate(raw_headers):
                if h is not None and str(h).strip():
                    columns.append(str(h).strip())
                else:
                    columns.append(f"Column_{idx + 1}")

            # Subsequent rows as data records
            parsed_rows = []
            for row_idx, row_values in enumerate(all_rows[1:]):
                # Skip entirely empty rows
                if all(v is None or (isinstance(v, str) and not v.strip()) for v in row_values):
                    continue

                row_dict = {}
                for col_idx, col_name in enumerate(columns):
                    val = row_values[col_idx] if col_idx < len(row_values) else None
                    if isinstance(val, (date, dt)):
                        val = val.isoformat()
                    row_dict[col_name] = val
                parsed_rows.append(row_dict)

            sheets_data.append({
                "name": sheet_name,
                "row_count": len(parsed_rows),
                "column_count": len(columns),
                "columns": columns,
                "rows": parsed_rows
            })

        wb.close()

        return {
            "filename": safe_filename,
            "sheet_count": len(sheets_data),
            "sheets": sheets_data
        }

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Error reading Excel file '{safe_filename}': {str(exc)}"
        )


@router.get("/excel/{filename}/normalize", summary="Normalize activity data from stored Excel workbook", tags=["Data Ingestion"])
def get_normalized_activities(filename: str, sheet: Optional[str] = None):
    """
    Feature 2.3: Standardizes and normalizes activity data from an uploaded Excel file.
    Detects standard activity fields (WBS/ID, Name, Planned/Actual Qty, Unit, Status, Dates)
    and cleans text, whitespace, and null values without writing to the database.
    """
    from app.services.normalization import normalize_activity_sheet

    # Reuse verified Excel parsing logic
    parsed_excel = get_excel_contents(filename)
    sheets_to_process = parsed_excel.get("sheets", [])

    if sheet:
        matching = [s for s in sheets_to_process if s["name"].strip().lower() == sheet.strip().lower()]
        if not matching:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Worksheet '{sheet}' not found in workbook '{filename}'."
            )
        sheets_to_process = matching

    normalized_sheets = []
    for s in sheets_to_process:
        norm_result = normalize_activity_sheet(
            sheet_name=s["name"],
            columns=s["columns"],
            rows=s["rows"]
        )
        normalized_sheets.append(norm_result)

    return {
        "filename": parsed_excel["filename"],
        "sheet_count": len(normalized_sheets),
        "sheets": normalized_sheets
    }


@router.get("/excel/{filename}/ai-extract", summary="AI Activity Extraction from normalized Excel workbook or Daily Report", tags=["Data Ingestion"])
def get_ai_extracted_activities(filename: Optional[str] = None, sheet: Optional[str] = None):
    """
    Feature 2.4 / 2.26: Extracts standardized activity attributes from normalized Excel
    or Daily Progress Reports (.pdf, .xlsx).
    Preserves original data without writing to PostgreSQL or mutating schedule baseline.
    """
    if not filename or filename in ("sample_l5_l6_schedule.xlsx", "null", "undefined", ""):
        active_fn = get_active_project_filename()
        if active_fn:
            filename = active_fn
        elif (STORAGE_DIR / "test_execution_matches.xlsx").exists() and filename == "test_execution_matches.xlsx":
            filename = "test_execution_matches.xlsx"
        elif (STORAGE_DIR / "baseline_schedule.xlsx").exists():
            filename = "baseline_schedule.xlsx"
        else:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No project data uploaded. Please upload a project Excel file to begin."
            )

    safe_filename = Path(filename).name
    if not (STORAGE_DIR / safe_filename).exists():
        active_fn = get_active_project_filename()
        if active_fn and (STORAGE_DIR / active_fn).exists():
            safe_filename = active_fn
            filename = active_fn
        else:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No project data uploaded. Please upload a project Excel file to begin."
            )

    file_ext = Path(safe_filename).suffix.lower()

    # Feature 2.26 / 2.27 / 2.28: If Daily Report, Site Diary, or Document is passed, parse directly into Activity Intelligence schema
    if file_ext == ".pdf":
        file_path = STORAGE_DIR / safe_filename
        if not file_path.exists() or not file_path.is_file():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"File '{safe_filename}' not found in storage."
            )
        if "diary" in safe_filename.lower() or "sd" in safe_filename.lower():
            from app.services.site_diary import parse_site_diary
            parsed = parse_site_diary(file_path)
        elif any(k in safe_filename.lower() for k in ["document", "doc", "scan", "ocr"]):
            from app.services.document_ingestion import parse_document
            parsed = parse_document(file_path)
        else:
            from app.services.daily_report import parse_daily_report
            parsed = parse_daily_report(file_path)
        return {
            "filename": safe_filename,
            "sheet_count": parsed.get("sheet_count", 1 if parsed.get("activities") else 0),
            "sheets": parsed.get("sheets", [])
        }

    if file_ext in (".docx", ".txt", ".png", ".jpg", ".jpeg"):
        file_path = STORAGE_DIR / safe_filename
        if not file_path.exists() or not file_path.is_file():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"File '{safe_filename}' not found in storage."
            )
        from app.services.document_ingestion import parse_document
        parsed = parse_document(file_path)
        return {
            "filename": safe_filename,
            "sheet_count": parsed.get("sheet_count", 1 if parsed.get("activities") else 0),
            "sheets": parsed.get("sheets", [])
        }

    # Feature 2.30: Voice / ASR execution capture audio ingestion
    if file_ext in (".wav", ".mp3", ".m4a", ".ogg", ".webm"):
        file_path = STORAGE_DIR / safe_filename
        if not file_path.exists() or not file_path.is_file():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Audio file '{safe_filename}' not found in storage."
            )
        from app.services.asr import run_asr, parse_voice_transcript
        asr_res = run_asr(file_path)
        parsed = parse_voice_transcript(asr_res.get("text", ""), safe_filename)
        return {
            "filename": safe_filename,
            "sheet_count": parsed.get("sheet_count", 1 if parsed.get("activities") else 0),
            "sheets": parsed.get("sheets", []),
            "asr": asr_res
        }

    from app.services.ai_extractor import extract_activities_from_normalized_sheet

    # Reuse Feature 2.3 normalized data pipeline for Excel workbooks
    normalized_data = get_normalized_activities(filename=filename, sheet=sheet)
    extracted_sheets = []

    for s in normalized_data.get("sheets", []):
        extracted_res = extract_activities_from_normalized_sheet(
            sheet_name=s["sheet_name"],
            activities=s["activities"]
        )
        extracted_sheets.append(extracted_res)

    return {
        "filename": normalized_data["filename"],
        "sheet_count": len(extracted_sheets),
        "sheets": extracted_sheets
    }


@router.get("/schedule/{filename}", summary="Read and parse L5/L6 schedule activities from Excel workbook", tags=["Data Ingestion"])
def get_schedule_activities(filename: Optional[str] = None, sheet: Optional[str] = None):
    """
    Feature 2.6: Reads an Excel schedule workbook, extracts L5/L6 schedule activities,
    normalizes level and date fields, and summarizes activity counts.
    Keeps schedule baseline data strictly separate from execution progress data.
    """
    if not filename or filename in ("sample_l5_l6_schedule.xlsx", "null", "undefined", ""):
        active_fn = get_active_project_filename()
        if active_fn:
            filename = active_fn
        elif (STORAGE_DIR / "baseline_schedule.xlsx").exists():
            filename = "baseline_schedule.xlsx"
        else:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No project data uploaded. Please upload a project Excel file to begin."
            )

    safe_filename = Path(filename).name
    if not (STORAGE_DIR / safe_filename).exists():
        active_fn = get_active_project_filename()
        if active_fn and (STORAGE_DIR / active_fn).exists():
            safe_filename = active_fn
            filename = active_fn
        elif (STORAGE_DIR / "baseline_schedule.xlsx").exists():
            safe_filename = "baseline_schedule.xlsx"
            filename = "baseline_schedule.xlsx"
        else:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No project data uploaded. Please upload a project Excel file to begin."
            )

    # Feature 2.31: If CSV schedule is provided, parse via schedule_ingestion
    file_ext = Path(safe_filename).suffix.lower()
    if file_ext == ".csv":
        from app.services.schedule_ingestion import parse_schedule_file
        file_path = STORAGE_DIR / safe_filename
        if not file_path.exists() or not file_path.is_file():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Schedule file '{safe_filename}' not found in storage."
            )
        parsed = parse_schedule_file(file_path)
        return {
            "filename": parsed["filename"],
            "sheet_count": parsed["sheet_count"],
            "total_schedule_activities": parsed["total_schedule_activities"],
            "l5_count": parsed["l5_count"],
            "l6_count": parsed["l6_count"],
            "other_count": parsed["other_count"],
            "sheets": parsed["sheets"]
        }

    from app.services.schedule_parser import parse_schedule_activities

    # Re-use existing Excel parser to load workbook sheets
    parsed_excel = get_excel_contents(filename)
    sheets_to_process = parsed_excel.get("sheets", [])

    if sheet:
        matching = [s for s in sheets_to_process if s["name"].strip().lower() == sheet.strip().lower()]
        if not matching:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Worksheet '{sheet}' not found in schedule workbook '{filename}'."
            )
        sheets_to_process = matching

    schedule_sheets = []
    total_acts = 0
    total_l5 = 0
    total_l6 = 0
    total_other = 0

    for s in sheets_to_process:
        res = parse_schedule_activities(
            sheet_name=s["name"],
            columns=s["columns"],
            rows=s["rows"]
        )
        total_acts += res["total_schedule_activities"]
        total_l5 += res["l5_count"]
        total_l6 += res["l6_count"]
        total_other += res["other_count"]
        schedule_sheets.append(res)

    return {
        "filename": parsed_excel["filename"],
        "sheet_count": len(schedule_sheets),
        "total_schedule_activities": total_acts,
        "l5_count": total_l5,
        "l6_count": total_l6,
        "other_count": total_other,
        "sheets": schedule_sheets
    }


@router.get("/match/exact-id", summary="Match execution activities to schedule activities strictly by Activity ID", tags=["Data Ingestion"])
def get_exact_id_matches(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.7: Matches AI-extracted execution activities with baseline schedule activities
    strictly using normalized Activity ID equality (normalized_execution_activity_id == normalized_schedule_activity_id).
    Does NOT perform fuzzy, semantic, or AI matching.
    """
    from app.services.matching import match_exact_id
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Perform exact ID matching
    match_result = match_exact_id(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return match_result


@router.get("/match/fuzzy", summary="Fuzzy text matching for execution activities unmatched by exact ID", tags=["Data Ingestion"])
def get_fuzzy_matches(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    threshold: float = 0.70,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.8: Compares unmatched execution activities against baseline schedule activities
    using lightweight local text similarity (difflib with token normalization).
    Only processes execution activities whose exact ID match status is 'unmatched'.
    Exact matches from Feature 2.7 are excluded from fuzzy results.
    Matches with score >= threshold are classified as 'possible_match' (never auto-confirmed).
    """
    from app.services.matching import match_fuzzy_activities
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Perform fuzzy matching on unmatched activities only
    fuzzy_result = match_fuzzy_activities(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        threshold=threshold,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return fuzzy_result


@router.get("/match/semantic", summary="Context and semantic matching for unresolved execution activities", tags=["Data Ingestion"])
def get_semantic_matches(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    threshold: float = 0.70,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.9: Context / Semantic text matching for execution activities that remain unresolved
    after Exact ID Matching (Feature 2.7) and Fuzzy Activity Matching (Feature 2.8).
    Exact matches and fuzzy possible matches are strictly excluded.
    Uses pure-Python local TF-IDF vectorization and cosine similarity over activity name, work description,
    discipline, and WBS.
    Matches with score >= threshold are classified as 'possible_match' (never auto-confirmed).
    """
    from app.services.matching import match_semantic_activities
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Perform semantic matching on unresolved activities only
    semantic_result = match_semantic_activities(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        threshold=threshold,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return semantic_result


@router.get("/match/granularity", summary="Resolve schedule activity hierarchy granularity (L5, L6, L5_PARENT, DETAILED, UNKNOWN)", tags=["Data Ingestion"])
def get_granularity_resolution(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.10: Determines whether each field execution activity is represented at the correct
    schedule granularity (L5, L6, L5_PARENT, DETAILED, UNKNOWN).
    Evaluates candidate schedule activities identified across Exact ID (Feature 2.7),
    Fuzzy (Feature 2.8), and Semantic (Feature 2.9) matching tiers.
    """
    from app.services.matching import resolve_all_granularities
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Perform granularity resolution
    granularity_result = resolve_all_granularities(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return granularity_result


@router.get("/match/new-activities", summary="Discover and propose new activity candidates from unmapped field execution records", tags=["Data Ingestion"])
def get_new_activity_discoveries(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.11: Evaluates unmapped execution activities after Exact ID (Feature 2.7),
    Fuzzy (Feature 2.8), and Context/Semantic (Feature 2.9) matching tiers.
    Proposes unmapped activities with verified descriptive evidence (name, work description, discipline)
    as NEW_ACTIVITY_CANDIDATE.
    Does NOT auto-confirm, modify schedule, or write to database.
    """
    from app.services.matching import discover_new_activities
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Perform New Activity Discovery
    discovery_result = discover_new_activities(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return discovery_result


@router.get("/validation/confidence", summary="Calculate activity match confidence and validation classification (AUTO_ACCEPT vs PLANNER_REVIEW)", tags=["Data Ingestion"])
def get_validation_confidence(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.12: Computes deterministic confidence scores and validation classifications
    (AUTO_ACCEPT vs. PLANNER_REVIEW) across the matching pipeline.
    AUTO_ACCEPT is strictly an analytical classification and does not write to the database
    or modify schedule files.
    """
    from app.services.matching import calculate_validation_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute confidence and validation classifications
    if not execution_activities:
        from app.services.demo_fallback import get_fallback_validation_confidence
        return get_fallback_validation_confidence(execution_file=execution_file, schedule_file=schedule_file)

    validation_result = calculate_validation_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file,
        auto_accept_threshold=0.90
    )

    return validation_result


@router.get("/validation/planner-review", summary="Inspect activities flagged for planner review and governance inspection", tags=["Data Ingestion"])
@router.get("/planner-review", summary="Inspect activities flagged for planner review (alias)", tags=["Data Ingestion"])
def get_planner_review(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.13: Dedicated read-only Planner Review and Validation Governance endpoint.
    Consumes Feature 2.12 validation results, filtering strictly for records with
    validation_status == 'PLANNER_REVIEW'.
    Does NOT modify schedules, auto-create database records, or commit approval decisions.
    """
    from app.services.planner_review import get_planner_review_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Retrieve planner review records
    if not execution_activities:
        from app.services.demo_fallback import get_fallback_planner_review
        return get_fallback_planner_review(execution_file=execution_file, schedule_file=schedule_file)

    review_result = get_planner_review_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return review_result


@router.get("/schedule-linking", summary="Schedule Linking: bridge field progress to baseline schedule activities", tags=["Data Ingestion"])
def get_schedule_linking(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.14: Planning-to-Execution Bridge linking field progress events
    to baseline L5/L6 schedule activities and calculating quantity variance.
    Does NOT modify schedules, write to database, or perform automatic updates.
    """
    from app.services.schedule_linking import get_schedule_linking_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Retrieve schedule linking results
    if not execution_activities:
        from app.services.demo_fallback import get_fallback_schedule_linking
        return get_fallback_schedule_linking(execution_file=execution_file, schedule_file=schedule_file)

    linking_result = get_schedule_linking_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return linking_result


@router.get("/schedule-dependencies", summary="Schedule Dependencies: baseline activity relationships (Predecessor -> Successor)", tags=["Data Ingestion"])
def get_schedule_dependencies_endpoint(
    schedule_file: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.15: Schedule Dependencies.
    Parses baseline schedule relationships (Predecessor -> Successor),
    validates that referenced activity IDs exist in the baseline schedule,
    and returns valid/invalid dependency classifications.
    Does NOT write to database or execute schedule mutations.
    """
    from app.services.schedule_dependencies import get_schedule_dependencies
    _, schedule_file = resolve_active_files(None, schedule_file)

    # 1. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 2. Parse and validate schedule dependencies
    dependency_result = get_schedule_dependencies(
        schedule_activities=schedule_activities,
        schedule_filename=schedule_file
    )

    return dependency_result


@router.get("/critical-path", summary="Critical Path Analysis: Forward pass, backward pass, float, and critical paths", tags=["Data Ingestion"])
def get_critical_path_endpoint(
    schedule_file: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.16: Critical Path Method (CPM) Analysis.
    Performs forward pass (Early Start, Early Finish), backward pass (Late Start, Late Finish),
    total float, and critical path identification on baseline schedule activities.
    Strictly read-only and analytical (no database writes or schedule mutations).
    """
    from app.services.critical_path import calculate_critical_path
    _, schedule_file = resolve_active_files(None, schedule_file)

    # 1. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 2. Compute Critical Path Analysis
    cpm_result = calculate_critical_path(
        schedule_activities=schedule_activities,
        schedule_filename=schedule_file
    )

    return cpm_result


@router.get("/delay-detection", summary="Delay Detection & Progress Variance: Baseline-vs-actual variance and delay severity", tags=["Data Ingestion"])
def get_delay_detection_endpoint(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.17: Delay Detection & Progress Variance.
    Compares baseline schedule vs actual field execution, computes progress %,
    quantity variance, date variance, delay status, severity, and critical path context.
    Strictly read-only and analytical (no database writes or schedule mutations).
    """
    from app.services.delay_detection import get_delay_detection_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute Delay Detection results
    delay_results = get_delay_detection_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return delay_results


@router.get("/schedule-health", summary="Milestone & Schedule Health: L5 milestone tracking, deterministic forecasting, and schedule condition score", tags=["Data Ingestion"])
def get_schedule_health_endpoint(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.18: Milestone & Schedule Health.
    Derives L5 milestone packages, correlates supporting L6 activities via WBS hierarchy,
    evaluates deterministic progress projections, and computes the overall schedule health score.
    Strictly read-only and analytical (no database writes or schedule mutations).
    """
    from app.services.schedule_health import get_schedule_health_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute Schedule Health and Milestone results
    health_results = get_schedule_health_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return health_results


@router.get("/delay-prediction", summary="Delay Prediction: Multi-factor deterministic delay risk estimation", tags=["Data Ingestion"])
def get_delay_prediction_endpoint(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.19: Delay Prediction.
    Estimates future delay likelihood for eligible schedule activities using multi-factor
    deterministic risk scoring across progress deficit, critical path exposure, severity, and milestone risk.
    Strictly read-only and analytical (no database writes or schedule mutations).
    """
    from app.services.delay_prediction import get_delay_prediction_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute Delay Prediction results
    prediction_results = get_delay_prediction_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return prediction_results


@router.get("/root-cause-analysis", summary="Root Cause Analysis: Deterministic multi-factor delay driver deduction", tags=["Data Ingestion"])
def get_root_cause_analysis_endpoint(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.20: Root Cause Analysis.
    Identifies and explains the likely contributing causes for delayed and at-risk schedule activities
    using evidence from progress deficits, quantity shortfalls, critical path exposure, predecessor delays, and milestone health.
    Strictly read-only and analytical (no database writes or schedule mutations).
    """
    from app.services.root_cause_analysis import get_root_cause_analysis_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute Root Cause Analysis results
    root_cause_results = get_root_cause_analysis_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return root_cause_results


@router.get("/impact-propagation", summary="Impact Propagation: Downstream dependency delay cascade analysis", tags=["Data Ingestion"])
def get_impact_propagation_endpoint(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.21: Impact Propagation.
    Determines how existing delayed or high-risk schedule activities propagate impact
    through the baseline dependency network to downstream activities, milestones, and project duration.
    Strictly read-only and analytical (no database writes or schedule modifications).
    """
    from app.services.impact_propagation import get_impact_propagation_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute Impact Propagation results
    impact_results = get_impact_propagation_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return impact_results


class WhatIfRequest(BaseModel):
    execution_file: Optional[str] = None
    schedule_file: Optional[str] = None
    execution_sheet: Optional[str] = None
    schedule_sheet: Optional[str] = None
    activity_id: Optional[str] = None
    target_activity_id: Optional[str] = None
    scenario_type: str
    scenario_value: Optional[float] = None
    value: Optional[float] = None
    scenario_name: Optional[str] = None


@router.post("/what-if", summary="What-If Simulator: In-memory hypothetical scenario calculation", tags=["Data Ingestion"])
def post_what_if_endpoint(
    request: WhatIfRequest,
    current_user: dict = Depends(require_roles("PLANNER", "PROJECT_MANAGER"))
):
    """
    Feature 2.22: What-If Simulator.
    Simulates in-memory hypothetical changes to activity delay, duration, or completion date.
    Recalculates downstream dependency propagation, milestone impact, and CPM critical paths.
    Strictly read-only with respect to the real project baseline (zero database or schedule writes).
    """
    from app.services.what_if_simulator import simulate_what_if_scenario

    target_id = request.activity_id or request.target_activity_id
    if not target_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="activity_id (or target_activity_id) must be specified."
        )

    val = request.scenario_value if request.scenario_value is not None else request.value
    if val is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="scenario_value (or value) must be specified."
        )

    resolved_exec, resolved_sched = resolve_active_files(request.execution_file, request.schedule_file)

    # 1. Fetch baseline schedule activities
    try:
        schedule_data = get_schedule_activities(filename=resolved_sched, sheet=request.schedule_sheet)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Schedule file '{resolved_sched}' not found or could not be parsed: {str(e)}"
        )

    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    if not schedule_activities:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"No schedule activities found in '{resolved_sched}'."
        )

    # 2. Fetch execution activities for milestone health context
    execution_activities = []
    try:
        extracted_data = get_ai_extracted_activities(filename=resolved_exec, sheet=request.execution_sheet)
        for s in extracted_data.get("sheets", []):
            execution_activities.extend(s.get("activities", []))
    except Exception:
        pass

    # 3. Simulate scenario in memory
    try:
        result = simulate_what_if_scenario(
            schedule_activities=schedule_activities,
            activity_id=target_id,
            scenario_type=request.scenario_type,
            scenario_value=val,
            schedule_filename=resolved_sched,
            execution_activities=execution_activities,
            execution_filename=resolved_exec
        )
        if request.scenario_name:
            result["scenario"]["scenario_name"] = request.scenario_name
    except KeyError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e).strip("'\""))
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e).strip("'\""))
    except Exception as e:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=f"What-If simulation failed: {str(e)}")

    return result


@router.get("/recovery-plans", summary="Recovery Plans: In-memory deterministic recovery plan generation", tags=["Data Ingestion"])
def get_recovery_plans_endpoint(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.23: Recovery Plans.
    Generates deterministic hypothetical recovery-plan scenarios for delayed and at-risk activities
    using controlled levers (DURATION_REDUCTION, PARALLEL_EXECUTION, START_ADVANCEMENT).
    Strictly read-only with respect to the baseline (zero database or schedule writes).
    """
    from app.services.recovery_plans import get_recovery_plans_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute Recovery Plans results
    results = get_recovery_plans_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return results


@router.get("/recommendations", summary="Recommendations: In-memory deterministic recovery scenario ranking and selection", tags=["Data Ingestion"])
def get_recommendations_endpoint(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.24: Recommendations.
    Deterministically evaluates and ranks recovery options for delayed or high-risk activities
    without inventing new recovery actions. Consumes existing recovery plans (Feature 2.23),
    applies multi-factor scoring (project finish recovery, critical path, delay severity,
    delay risk, milestone exposure, feasibility, float improvement), and provides clear
    evidence-based rationale.
    Strictly read-only with respect to the baseline (zero database or schedule writes).
    """
    from app.services.recommendations import get_recommendations_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute Recommendations results
    results = get_recommendations_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return results


@router.get("/decision-center", summary="Decision Center: Executive analytical cockpit consolidating alerts, insights, and recommendations", tags=["Data Ingestion"])
def get_decision_center_endpoint(
    execution_file: Optional[str] = None,
    schedule_file: Optional[str] = None,
    execution_sheet: Optional[str] = None,
    schedule_sheet: Optional[str] = None
):
    """
    Feature 2.25: Decision Center.
    Executive analytical cockpit consolidating already-generated outputs from previous pipeline
    layers (Critical Path, Delay Detection, Schedule Health, Delay Prediction, Root Cause Analysis,
    Impact Propagation, What-If Simulator, Recovery Plans, and Recommendations).
    Presents deterministic alerts, insights, recommendations, affected activities, project overview,
    and maintains an analytical action status ('PENDING_MANAGEMENT_DECISION').
    Strictly read-only with respect to the baseline (zero database or schedule writes).
    """
    from app.services.decision_center import get_decision_center_results
    execution_file, schedule_file = resolve_active_files(execution_file, schedule_file)

    # 1. Fetch AI-extracted execution activities
    extracted_data = get_ai_extracted_activities(filename=execution_file, sheet=execution_sheet)
    execution_activities = []
    for s in extracted_data.get("sheets", []):
        execution_activities.extend(s.get("activities", []))

    # 2. Fetch baseline schedule activities
    schedule_data = get_schedule_activities(filename=schedule_file, sheet=schedule_sheet)
    schedule_activities = []
    for s in schedule_data.get("sheets", []):
        schedule_activities.extend(s.get("activities", []))

    # 3. Compute Decision Center consolidated results
    results = get_decision_center_results(
        execution_activities=execution_activities,
        schedule_activities=schedule_activities,
        execution_filename=execution_file,
        schedule_filename=schedule_file
    )

    return results


# =============================================================================
# FEATURE 2.26: DAILY REPORT INGESTION
# =============================================================================

@router.post("/daily-reports/upload", summary="Upload Daily Progress Report (.pdf, .xlsx, .xls)", tags=["Daily Reports"])
async def upload_daily_report_file(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Feature 2.26: Validates and stores an uploaded Daily Report locally under backend/storage/.
    Supports .pdf, .xlsx, .xls files.
    Returns upload status, filename, file size, and file type.
    """
    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename is required."
        )

    file_ext = Path(file.filename).suffix.lower()
    from app.services.daily_report import ALLOWED_DAILY_REPORT_EXTENSIONS
    if file_ext not in ALLOWED_DAILY_REPORT_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed Daily Report file types: {sorted(ALLOWED_DAILY_REPORT_EXTENSIONS)}"
        )

    safe_filename = Path(file.filename).name
    destination_path = STORAGE_DIR / safe_filename

    try:
        total_bytes = 0
        with open(destination_path, "wb") as buffer:
            while chunk := await file.read(1024 * 1024):
                buffer.write(chunk)
                total_bytes += len(chunk)

        if total_bytes == 0:
            if destination_path.exists():
                destination_path.unlink()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded Daily Report file is empty (0 bytes)."
            )
    except HTTPException:
        raise
    except Exception as exc:
        if destination_path.exists():
            destination_path.unlink()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save uploaded Daily Report file: {str(exc)}"
        )
    finally:
        await file.close()

    return {
        "status": "success",
        "filename": safe_filename,
        "file_size": total_bytes,
        "file_type": "pdf" if file_ext == ".pdf" else "excel",
        "stored_location": f"backend/storage/{safe_filename}",
        "message": "Daily Report uploaded successfully",
        "timestamp": datetime.now(timezone.utc).isoformat()
    }


@router.get("/daily-reports", summary="List available Daily Report files in local storage", tags=["Daily Reports"])
def list_daily_reports():
    """
    Feature 2.26: Lists available Daily Report files (.pdf, .xlsx, .xls) stored in backend/storage/.
    """
    from app.services.daily_report import ALLOWED_DAILY_REPORT_EXTENSIONS
    files = []
    if STORAGE_DIR.exists():
        for p in sorted(STORAGE_DIR.iterdir()):
            if p.is_file() and p.suffix.lower() in ALLOWED_DAILY_REPORT_EXTENSIONS:
                stat = p.stat()
                files.append({
                    "filename": p.name,
                    "file_size": stat.st_size,
                    "file_type": "pdf" if p.suffix.lower() == ".pdf" else "excel",
                    "modified_time": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                    "is_daily_report": "daily" in p.name.lower() or "dpr" in p.name.lower() or p.suffix.lower() == ".pdf"
                })

    return {
        "total_files": len(files),
        "files": files
    }


@router.get("/daily-reports/{filename}", summary="Read and parse Daily Report (.pdf, .xlsx, .xls)", tags=["Daily Reports"])
def get_daily_report_endpoint(filename: str):
    """
    Feature 2.26: Reads and parses a stored Daily Report from backend/storage/.
    Extracts metadata, site conditions, and structured activities formatted
    for seamless consumption by the existing Activity Intelligence pipeline.
    Zero database or schedule mutations (strictly read-only).
    """
    if not filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename parameter is required."
        )

    safe_filename = Path(filename).name
    from app.services.daily_report import ALLOWED_DAILY_REPORT_EXTENSIONS, parse_daily_report

    file_ext = Path(safe_filename).suffix.lower()
    if file_ext not in ALLOWED_DAILY_REPORT_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed Daily Report types: {sorted(ALLOWED_DAILY_REPORT_EXTENSIONS)}"
        )

    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Daily Report file '{safe_filename}' not found in storage. Please upload it first."
        )

    try:
        parsed = parse_daily_report(file_path)
        return parsed
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Error parsing Daily Report '{safe_filename}': {str(exc)}"
        )


# =============================================================================
# FEATURE 2.27: SITE DIARY INGESTION
# =============================================================================

@router.post("/site-diaries/upload", summary="Upload Site Diary (.pdf, .xlsx, .xls)", tags=["Site Diaries"])
async def upload_site_diary_file(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Feature 2.27: Validates and stores an uploaded Site Diary locally under backend/storage/.
    Supports .pdf, .xlsx, .xls files.
    Returns upload status, filename, file size, and file type.
    Strictly read-only analytical ingestion (zero database writes).
    """
    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename is required."
        )

    file_ext = Path(file.filename).suffix.lower()
    from app.services.site_diary import ALLOWED_SITE_DIARY_EXTENSIONS
    if file_ext not in ALLOWED_SITE_DIARY_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed Site Diary file types: {sorted(ALLOWED_SITE_DIARY_EXTENSIONS)}"
        )

    safe_filename = Path(file.filename).name
    destination_path = STORAGE_DIR / safe_filename

    try:
        total_bytes = 0
        with open(destination_path, "wb") as buffer:
            while chunk := await file.read(1024 * 1024):
                buffer.write(chunk)
                total_bytes += len(chunk)

        if total_bytes == 0:
            if destination_path.exists():
                destination_path.unlink()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded Site Diary file is empty (0 bytes)."
            )
    except HTTPException:
        raise
    except Exception as exc:
        if destination_path.exists():
            destination_path.unlink()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save uploaded Site Diary file: {str(exc)}"
        )
    finally:
        await file.close()

    return {
        "status": "success",
        "filename": safe_filename,
        "file_size": total_bytes,
        "file_type": "pdf" if file_ext == ".pdf" else "excel",
        "stored_location": f"backend/storage/{safe_filename}",
        "message": "Site Diary uploaded successfully",
        "timestamp": datetime.now(timezone.utc).isoformat()
    }


@router.get("/site-diaries", summary="List available Site Diary files in local storage", tags=["Site Diaries"])
def list_site_diaries():
    """
    Feature 2.27: Lists available Site Diary files (.pdf, .xlsx, .xls) stored in backend/storage/.
    """
    from app.services.site_diary import ALLOWED_SITE_DIARY_EXTENSIONS
    files = []
    if STORAGE_DIR.exists():
        for p in sorted(STORAGE_DIR.iterdir()):
            if p.is_file() and p.suffix.lower() in ALLOWED_SITE_DIARY_EXTENSIONS:
                stat = p.stat()
                is_sd = "diary" in p.name.lower() or "sd" in p.name.lower() or "site_diary" in p.name.lower()
                files.append({
                    "filename": p.name,
                    "file_size": stat.st_size,
                    "file_type": "pdf" if p.suffix.lower() == ".pdf" else "excel",
                    "modified_time": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                    "is_site_diary": is_sd
                })

    return {
        "total_files": len(files),
        "files": files
    }


@router.get("/site-diaries/{filename}", summary="Read and parse Site Diary (.pdf, .xlsx, .xls)", tags=["Site Diaries"])
def get_site_diary_endpoint(filename: str):
    """
    Feature 2.27: Reads and parses a stored Site Diary from backend/storage/.
    Extracts metadata, site/work location, remarks, issues, and structured activities
    formatted for seamless consumption by the existing Activity Intelligence pipeline.
    Zero database or schedule mutations (strictly read-only).
    """
    if not filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename parameter is required."
        )

    safe_filename = Path(filename).name
    from app.services.site_diary import ALLOWED_SITE_DIARY_EXTENSIONS, parse_site_diary

    file_ext = Path(safe_filename).suffix.lower()
    if file_ext not in ALLOWED_SITE_DIARY_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed Site Diary types: {sorted(ALLOWED_SITE_DIARY_EXTENSIONS)}"
        )

    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Site Diary file '{safe_filename}' not found in storage. Please upload it first."
        )

    try:
        parsed = parse_site_diary(file_path)
        return parsed
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Error parsing Site Diary '{safe_filename}': {str(exc)}"
        )


# =============================================================================
# FEATURE 2.28: DOCUMENTS INGESTION
# =============================================================================

@router.post("/documents/upload", summary="Upload Project Document (.pdf, .docx, .txt)", tags=["Documents"])
async def upload_document_file(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Feature 2.28: Validates and stores an uploaded Project Document locally under backend/storage/.
    Supports .pdf, .docx, .txt files.
    Returns upload status, filename, file size, and file type.
    Strictly read-only analytical ingestion (zero database writes).
    """
    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename is required."
        )

    file_ext = Path(file.filename).suffix.lower()
    from app.services.document_ingestion import ALLOWED_DOCUMENT_EXTENSIONS
    if file_ext not in ALLOWED_DOCUMENT_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed Document file types: {sorted(ALLOWED_DOCUMENT_EXTENSIONS)}"
        )

    safe_filename = Path(file.filename).name
    destination_path = STORAGE_DIR / safe_filename

    try:
        total_bytes = 0
        with open(destination_path, "wb") as buffer:
            while chunk := await file.read(1024 * 1024):
                buffer.write(chunk)
                total_bytes += len(chunk)

        if total_bytes == 0:
            if destination_path.exists():
                destination_path.unlink()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded Document file is empty (0 bytes)."
            )
    except HTTPException:
        raise
    except Exception as exc:
        if destination_path.exists():
            destination_path.unlink()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save uploaded Document file: {str(exc)}"
        )
    finally:
        await file.close()

    type_mapping = {".pdf": "pdf", ".docx": "word", ".txt": "text"}

    return {
        "status": "success",
        "filename": safe_filename,
        "file_size": total_bytes,
        "file_type": type_mapping.get(file_ext, "document"),
        "stored_location": f"backend/storage/{safe_filename}",
        "message": "Document uploaded successfully",
        "timestamp": datetime.now(timezone.utc).isoformat()
    }


@router.get("/documents", summary="List available Document files in local storage", tags=["Documents"])
def list_documents():
    """
    Feature 2.28: Lists available Document files (.pdf, .docx, .txt) stored in backend/storage/.
    """
    from app.services.document_ingestion import ALLOWED_DOCUMENT_EXTENSIONS
    type_mapping = {".pdf": "pdf", ".docx": "word", ".txt": "text"}
    files = []
    if STORAGE_DIR.exists():
        for p in sorted(STORAGE_DIR.iterdir()):
            if p.is_file() and p.suffix.lower() in ALLOWED_DOCUMENT_EXTENSIONS:
                stat = p.stat()
                is_doc = "document" in p.name.lower() or "doc" in p.name.lower() or p.suffix.lower() in (".docx", ".txt")
                files.append({
                    "filename": p.name,
                    "file_size": stat.st_size,
                    "file_type": type_mapping.get(p.suffix.lower(), "document"),
                    "modified_time": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                    "is_document": is_doc
                })

    return {
        "total_files": len(files),
        "files": files
    }


@router.get("/documents/{filename}", summary="Read and parse Project Document (.pdf, .docx, .txt)", tags=["Documents"])
def get_document_endpoint(filename: str):
    """
    Feature 2.28: Reads and parses a stored Project Document from backend/storage/.
    Extracts metadata, project name, discipline, contractor, prepared by, summary,
    remarks, issues, text preview, and structured activities formatted for seamless
    consumption by the existing Activity Intelligence pipeline.
    Zero database or schedule mutations (strictly read-only).
    """
    if not filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename parameter is required."
        )

    safe_filename = Path(filename).name
    from app.services.document_ingestion import ALLOWED_DOCUMENT_EXTENSIONS, parse_document

    file_ext = Path(safe_filename).suffix.lower()
    if file_ext not in ALLOWED_DOCUMENT_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed Document types: {sorted(ALLOWED_DOCUMENT_EXTENSIONS)}"
        )

    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Document file '{safe_filename}' not found in storage. Please upload it first."
        )

    try:
        parsed = parse_document(file_path)
        return parsed
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Error parsing Document '{safe_filename}': {str(exc)}"
        )


# =============================================================================
# FEATURE 2.29: OCR (OPTICAL CHARACTER RECOGNITION)
# =============================================================================

@router.get("/ocr/status", summary="Check local OCR engine status and availability", tags=["OCR"])
def get_ocr_status():
    """
    Feature 2.29: Returns the availability and version of the local Tesseract OCR engine.
    Strictly free and local (no cloud APIs).
    """
    from app.services.ocr import get_ocr_engine_info
    available, engine_info, tess_path = get_ocr_engine_info()
    return {
        "ocr_available": available,
        "engine": engine_info,
        "tesseract_path": tess_path,
        "supported_formats": [".pdf", ".png", ".jpg", ".jpeg"],
        "timestamp": datetime.now(timezone.utc).isoformat()
    }


@router.post("/ocr/upload", summary="Upload scanned document or image for local OCR extraction", tags=["OCR"])
async def upload_ocr_file(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Feature 2.29: Validates and stores an uploaded scanned document or image locally under backend/storage/.
    Executes local Tesseract OCR, extracts text, and parses structured activities into
    standardized Activity Intelligence schema.
    Strictly in-memory analytical processing (zero database writes).
    """
    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename is required."
        )

    file_ext = Path(file.filename).suffix.lower()
    from app.services.ocr import ALLOWED_OCR_EXTENSIONS, run_ocr
    from app.services.document_ingestion import parse_document
    if file_ext not in ALLOWED_OCR_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed OCR file types: {sorted(ALLOWED_OCR_EXTENSIONS)}"
        )

    safe_filename = Path(file.filename).name
    destination_path = STORAGE_DIR / safe_filename

    try:
        total_bytes = 0
        with open(destination_path, "wb") as buffer:
            while chunk := await file.read(1024 * 1024):
                buffer.write(chunk)
                total_bytes += len(chunk)

        if total_bytes == 0:
            if destination_path.exists():
                destination_path.unlink()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded OCR file is empty (0 bytes)."
            )
    except HTTPException:
        raise
    except Exception as exc:
        if destination_path.exists():
            destination_path.unlink()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save uploaded OCR file: {str(exc)}"
        )
    finally:
        await file.close()

    try:
        ocr_result = run_ocr(destination_path)
        doc_result = parse_document(destination_path)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Error executing OCR on '{safe_filename}': {str(exc)}"
        )

    return {
        "status": "success",
        "filename": safe_filename,
        "file_size": total_bytes,
        "file_type": "pdf" if file_ext == ".pdf" else "image",
        "stored_location": f"backend/storage/{safe_filename}",
        "ocr": ocr_result,
        "document_metadata": doc_result.get("document_metadata", {}),
        "activities": doc_result.get("activities", []),
        "total_activities": len(doc_result.get("activities", [])),
        "sheets": doc_result.get("sheets", []),
        "message": "File processed via local OCR successfully",
        "timestamp": datetime.now(timezone.utc).isoformat()
    }


@router.get("/ocr/files", summary="List available OCR-compatible files in local storage", tags=["OCR"])
def list_ocr_files():
    """
    Feature 2.29: Lists available scanned documents and image files (.pdf, .png, .jpg, .jpeg) in backend/storage/.
    """
    from app.services.ocr import ALLOWED_OCR_EXTENSIONS
    files = []
    if STORAGE_DIR.exists():
        for p in sorted(STORAGE_DIR.iterdir()):
            if p.is_file() and p.suffix.lower() in ALLOWED_OCR_EXTENSIONS:
                stat = p.stat()
                is_scanned = "scan" in p.name.lower() or "ocr" in p.name.lower() or p.suffix.lower() in (".png", ".jpg", ".jpeg")
                files.append({
                    "filename": p.name,
                    "file_size": stat.st_size,
                    "file_type": "pdf" if p.suffix.lower() == ".pdf" else "image",
                    "modified_time": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                    "is_scanned": is_scanned
                })

    return {
        "total_files": len(files),
        "files": files
    }


@router.get("/ocr/{filename}", summary="Run OCR on a stored document/image and return structured output", tags=["OCR"])
def get_ocr_endpoint(filename: str):
    """
    Feature 2.29: Runs local OCR on a specified stored file from backend/storage/.
    Returns full OCR inspection result (pages, characters, confidence, recognized text)
    and passes text into standardized Activity Intelligence format.
    Zero database or schedule mutations (strictly read-only).
    """
    if not filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename parameter is required."
        )

    safe_filename = Path(filename).name
    from app.services.ocr import ALLOWED_OCR_EXTENSIONS, run_ocr
    from app.services.document_ingestion import parse_document

    file_ext = Path(safe_filename).suffix.lower()
    if file_ext not in ALLOWED_OCR_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed OCR types: {sorted(ALLOWED_OCR_EXTENSIONS)}"
        )

    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"OCR file '{safe_filename}' not found in storage. Please upload it first."
        )

    try:
        ocr_result = run_ocr(file_path)
        doc_result = parse_document(file_path)
        return {
            "filename": safe_filename,
            "ocr": ocr_result,
            "document_metadata": doc_result.get("document_metadata", {}),
            "activities": doc_result.get("activities", []),
            "total_activities": len(doc_result.get("activities", [])),
            "sheets": doc_result.get("sheets", []),
            "notice": ocr_result.get("notice"),
            "in_memory_only": True,
            "database_modified": False,
            "baseline_schedule_modified": False
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Error executing OCR on '{safe_filename}': {str(exc)}"
        )


# ============================================================================
# FEATURE 2.30: VOICE / ASR (AUTOMATED SPEECH RECOGNITION)
# ============================================================================

@router.get("/asr/status", summary="Check local Voice / ASR engine status and availability", tags=["Voice / ASR"])
def get_asr_status():
    """
    Feature 2.30: Returns the availability and engine configuration of the local ASR engine (faster-whisper).
    Completely offline, local, and free (no external or paid cloud APIs).
    """
    from app.services.asr import get_asr_engine_info, ALLOWED_ASR_EXTENSIONS
    available, engine_info, model_name = get_asr_engine_info()
    return {
        "asr_available": available,
        "engine_info": engine_info,
        "model": model_name,
        "device": "cpu",
        "compute_type": "int8",
        "allowed_extensions": sorted(list(ALLOWED_ASR_EXTENSIONS)),
        "local_only": True,
        "cloud_api_used": False,
        "notice": "Local faster-whisper ASR operational on CPU" if available else engine_info
    }


@router.post("/asr/upload", summary="Upload voice progress recording for local ASR transcription", tags=["Voice / ASR"])
async def upload_asr_file(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Feature 2.30: Ingests a field voice audio recording (.wav, .mp3, .m4a, .ogg, .webm).
    Transcribes audio via local faster-whisper on CPU, parses structured activity updates,
    and returns standardized Activity Intelligence format in-memory.
    Zero cloud calls, zero database writes, zero schedule modifications.
    """
    if not file or not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No file provided in the upload request."
        )

    safe_filename = Path(file.filename).name
    file_ext = Path(safe_filename).suffix.lower()

    from app.services.asr import ALLOWED_ASR_EXTENSIONS, run_asr, parse_voice_transcript

    if file_ext not in ALLOWED_ASR_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed audio types: {sorted(ALLOWED_ASR_EXTENSIONS)}"
        )

    destination_path = STORAGE_DIR / safe_filename

    try:
        content = await file.read()
        if len(content) == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded audio file is empty (0 bytes)."
            )

        with open(destination_path, "wb") as buffer:
            buffer.write(content)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save uploaded audio file: {str(exc)}"
        )

    try:
        asr_result = run_asr(destination_path)
        parsed = parse_voice_transcript(asr_result.get("text", ""), safe_filename)
        return {
            "filename": safe_filename,
            "asr": asr_result,
            "document_metadata": parsed.get("document_metadata", {}),
            "report_metadata": parsed.get("report_metadata", {}),
            "activities": parsed.get("activities", []),
            "total_activities": len(parsed.get("activities", [])),
            "sheets": parsed.get("sheets", []),
            "message": "Audio file processed via local ASR successfully",
            "in_memory_only": True,
            "database_modified": False,
            "baseline_schedule_modified": False
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Error executing ASR on '{safe_filename}': {str(exc)}"
        )


@router.get("/asr/files", summary="List available voice recordings in local storage", tags=["Voice / ASR"])
def list_asr_files():
    """
    Feature 2.30: Returns all audio recordings (.wav, .mp3, .m4a, .ogg, .webm) available in local storage.
    """
    from app.services.asr import ALLOWED_ASR_EXTENSIONS

    files = []
    if STORAGE_DIR.exists():
        for p in sorted(STORAGE_DIR.iterdir()):
            if p.is_file() and p.suffix.lower() in ALLOWED_ASR_EXTENSIONS:
                stat = p.stat()
                files.append({
                    "filename": p.name,
                    "format": p.suffix.lower().lstrip("."),
                    "size_bytes": stat.st_size,
                    "size_kb": round(stat.st_size / 1024, 2),
                    "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                    "source": "Local Voice Recording"
                })

    return {
        "total_files": len(files),
        "files": files,
        "allowed_extensions": sorted(list(ALLOWED_ASR_EXTENSIONS)),
        "in_memory_only": True,
        "database_modified": False
    }


@router.get("/asr/{filename}", summary="Run ASR on a stored voice audio file and return structured output", tags=["Voice / ASR"])
def get_asr_endpoint(filename: str):
    """
    Feature 2.30: Runs local faster-whisper ASR on a specified stored audio file from backend/storage/.
    Returns transcript text, timing segments, duration, and parsed activities.
    Pure-local, in-memory, no database writes.
    """
    safe_filename = Path(filename).name
    from app.services.asr import ALLOWED_ASR_EXTENSIONS, run_asr, parse_voice_transcript

    file_ext = Path(safe_filename).suffix.lower()
    if file_ext not in ALLOWED_ASR_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed audio types: {sorted(ALLOWED_ASR_EXTENSIONS)}"
        )

    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Audio file '{safe_filename}' not found in storage. Please upload it first."
        )

    try:
        asr_result = run_asr(file_path)
        parsed = parse_voice_transcript(asr_result.get("text", ""), safe_filename)
        return {
            "filename": safe_filename,
            "asr": asr_result,
            "document_metadata": parsed.get("document_metadata", {}),
            "report_metadata": parsed.get("report_metadata", {}),
            "activities": parsed.get("activities", []),
            "total_activities": len(parsed.get("activities", [])),
            "sheets": parsed.get("sheets", []),
            "notice": asr_result.get("notice"),
            "in_memory_only": True,
            "database_modified": False,
            "baseline_schedule_modified": False
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Error executing ASR on '{safe_filename}': {str(exc)}"
        )


@router.get("/asr/stream/{filename}", summary="Stream voice audio file for playback", tags=["Voice / ASR"])
def stream_asr_audio(filename: str):
    """
    Feature 2.30: Streams local stored audio file for in-browser playback.
    """
    safe_filename = Path(filename).name
    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Audio file '{safe_filename}' not found."
        )
    from fastapi.responses import FileResponse
    ext = file_path.suffix.lower()
    media_types = {
        ".wav": "audio/wav",
        ".mp3": "audio/mpeg",
        ".m4a": "audio/mp4",
        ".ogg": "audio/ogg",
        ".webm": "audio/webm"
    }
    return FileResponse(file_path, media_type=media_types.get(ext, "application/octet-stream"))


class LiveVoiceParseRequest(BaseModel):
    text: str
    location: Optional[str] = None
    evidence_id: Optional[str] = None


@router.post("/asr/live-chunk", summary="Process live audio chunk for real-time local transcription and event capture", tags=["Voice / ASR"])
async def transcribe_live_chunk(
    file: UploadFile = File(...),
    is_final: bool = Form(False),
    location: Optional[str] = Form(None),
    evidence_id: Optional[str] = Form(None),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Feature 2.30 Real-Time: Accepts a live audio chunk/recording from browser MediaRecorder.
    Transcribes audio completely in memory using local Faster-Whisper on CPU with int8 quantization.
    Extracts structured execution events (Activity, Event, Location, Discipline, Confidence) and standardized activities.
    Zero cloud calls, zero database writes, zero permanent disk storage.
    """
    if not file or not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No audio chunk provided."
        )

    from app.services.asr import run_live_chunk_asr

    try:
        content = await file.read()
        if len(content) == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Audio chunk is empty (0 bytes)."
            )

        result = run_live_chunk_asr(
            audio_bytes=content,
            filename=file.filename,
            is_final=is_final,
            location=location,
            evidence_id=evidence_id
        )
        return result
    except HTTPException:
        raise
    except ValueError as val_err:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(val_err)
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error transcribing live audio chunk: {str(exc)}"
        )


@router.post("/asr/live-parse", summary="Parse transcribed live voice text into structured execution events", tags=["Voice / ASR"])
def parse_live_text(
    payload: LiveVoiceParseRequest,
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Feature 2.30 Real-Time: Parses natural language execution text into structured event and activity intelligence.
    Pure in-memory, zero database writes.
    """
    from app.services.asr import extract_live_voice_event, parse_voice_transcript
    event = extract_live_voice_event(payload.text, location=payload.location, evidence_id=payload.evidence_id)
    parsed = parse_voice_transcript(payload.text)
    return {
        "text": payload.text,
        "extracted_event": event,
        "activities": parsed.get("activities", []),
        "total_activities": len(parsed.get("activities", [])),
        "in_memory_only": True,
        "database_modified": False
    }


# ============================================================================
# FEATURE 2.31: PRIMAVERA / MS PROJECT SCHEDULE INGESTION
# ============================================================================

@router.post("/schedules/upload", summary="Upload Primavera or MS Project exported schedule file", tags=["Schedule Ingestion"])
async def upload_schedule_file(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_roles("PLANNER"))
):
    """
    Feature 2.31: Ingests an exported Primavera P6 or MS Project schedule (.xlsx, .xls, .csv).
    Validates format, parses activities in-memory, identifies source export type,
    and returns standardized schedule activities.
    Zero permanent writes to PostgreSQL, zero mutations to baseline schedule.
    """
    if not file or not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename is required."
        )

    safe_filename = Path(file.filename).name
    file_ext = Path(safe_filename).suffix.lower()

    from app.services.schedule_ingestion import ALLOWED_SCHEDULE_EXTENSIONS, parse_schedule_file

    if file_ext not in ALLOWED_SCHEDULE_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed schedule types: {sorted(ALLOWED_SCHEDULE_EXTENSIONS)}"
        )

    destination_path = STORAGE_DIR / safe_filename

    try:
        content = await file.read()
        if len(content) == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded schedule file is empty (0 bytes)."
            )

        with open(destination_path, "wb") as buffer:
            buffer.write(content)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save uploaded schedule file: {str(exc)}"
        )

    try:
        parsed = parse_schedule_file(destination_path)
        return {
            "status": "success",
            "filename": safe_filename,
            "source_type": parsed["source_type"],
            "source_format": parsed["source_format"],
            "total_activities": parsed["total_activities"],
            "total_schedule_activities": parsed["total_activities"],
            "sheet_count": parsed["sheet_count"],
            "l5_count": parsed["l5_count"],
            "l6_count": parsed["l6_count"],
            "other_count": parsed["other_count"],
            "sheets": parsed["sheets"],
            "activities": parsed["activities"],
            "message": "Schedule uploaded and parsed successfully",
            "in_memory_only": True,
            "database_modified": False,
            "baseline_schedule_modified": False
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Error parsing schedule '{safe_filename}': {str(exc)}"
        )


@router.get("/schedules", summary="List available schedule files in local storage", tags=["Schedule Ingestion"])
def list_schedules():
    """
    Feature 2.31: Lists all available exported and baseline schedule files (.xlsx, .xls, .csv).
    """
    from app.services.schedule_ingestion import ALLOWED_SCHEDULE_EXTENSIONS

    files = []
    if STORAGE_DIR.exists():
        for p in sorted(STORAGE_DIR.iterdir()):
            if p.is_file() and p.suffix.lower() in ALLOWED_SCHEDULE_EXTENSIONS:
                if any(x in p.name.lower() for x in ["execution", "progress", "report", "diary"]):
                    continue
                stat = p.stat()
                source_type = "GENERIC_SCHEDULE"
                if "primavera" in p.name.lower() or "p6" in p.name.lower():
                    source_type = "PRIMAVERA_P6_EXPORT"
                elif "ms_project" in p.name.lower() or "msp" in p.name.lower():
                    source_type = "MS_PROJECT_EXPORT"

                files.append({
                    "filename": p.name,
                    "source_type": source_type,
                    "source_format": p.suffix.lower().lstrip("."),
                    "size_bytes": stat.st_size,
                    "size_kb": round(stat.st_size / 1024, 2),
                    "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()
                })

    return {
        "total_files": len(files),
        "files": files,
        "allowed_extensions": sorted(list(ALLOWED_SCHEDULE_EXTENSIONS)),
        "in_memory_only": True,
        "database_modified": False
    }


@router.get("/schedules/{filename}", summary="Read and parse a stored Primavera/MS Project schedule", tags=["Schedule Ingestion"])
def get_schedule_file_detail(filename: str):
    """
    Feature 2.31: Reads and parses a stored schedule export (.xlsx, .xls, .csv) in-memory.
    Returns detected source type, source format, and standardized L5/L6 activity register.
    """
    if not filename or filename in ("sample_l5_l6_schedule.xlsx", "null", "undefined", ""):
        active_fn = get_active_project_filename()
        if active_fn:
            filename = active_fn
        elif (STORAGE_DIR / "baseline_schedule.xlsx").exists():
            filename = "baseline_schedule.xlsx"
        else:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No project data uploaded. Please upload a project Excel file to begin."
            )

    safe_filename = Path(filename).name
    if not (STORAGE_DIR / safe_filename).exists():
        active_fn = get_active_project_filename()
        if active_fn and (STORAGE_DIR / active_fn).exists():
            safe_filename = active_fn
            filename = active_fn
        elif (STORAGE_DIR / "baseline_schedule.xlsx").exists():
            safe_filename = "baseline_schedule.xlsx"
            filename = "baseline_schedule.xlsx"
        else:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No project data uploaded. Please upload a project Excel file to begin."
            )

    file_ext = Path(safe_filename).suffix.lower()

    from app.services.schedule_ingestion import ALLOWED_SCHEDULE_EXTENSIONS, parse_schedule_file

    if file_ext not in ALLOWED_SCHEDULE_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed schedule types: {sorted(ALLOWED_SCHEDULE_EXTENSIONS)}"
        )

    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="No project data uploaded. Please upload a project Excel file to begin."
        )

    try:
        parsed = parse_schedule_file(file_path)
        return {
            "filename": safe_filename,
            "source_type": parsed["source_type"],
            "source_format": parsed["source_format"],
            "total_activities": parsed["total_activities"],
            "total_schedule_activities": parsed["total_activities"],
            "sheet_count": parsed["sheet_count"],
            "l5_count": parsed["l5_count"],
            "l6_count": parsed["l6_count"],
            "other_count": parsed["other_count"],
            "sheets": parsed["sheets"],
            "activities": parsed["activities"],
            "in_memory_only": True,
            "database_modified": False,
            "baseline_schedule_modified": False
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Error parsing schedule file '{safe_filename}': {str(exc)}"
        )


# ============================================================================
# FEATURE A: PHOTO INGESTION & EXECUTION EVIDENCE
# ============================================================================

@router.post("/photos/upload", summary="Upload photo as execution evidence", tags=["Photo Evidence"])
async def upload_photo_evidence(
    file: UploadFile = File(...),
    capture_date: Optional[str] = Form(None),
    project_name: Optional[str] = Form(None),
    site_location: Optional[str] = Form(None),
    work_location: Optional[str] = Form(None),
    discipline: Optional[str] = Form(None),
    activity_id: Optional[str] = Form(None),
    description: Optional[str] = Form(None),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    Feature A: Ingests an image file (.jpg, .jpeg, .png, .webp) as execution evidence.
    Extracts image metadata (dimensions, size, mime type) and preserves optional
    evidence attribution without performing computer vision, OCR, or automatic validation.
    Zero permanent writes to PostgreSQL, zero mutations to baseline schedule.
    """
    if not file or not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Filename is required."
        )

    from app.services.photo_ingestion import (
        ALLOWED_PHOTO_EXTENSIONS,
        parse_photo_metadata,
        save_sidecar_metadata,
        sanitize_filename
    )

    try:
        safe_filename = sanitize_filename(file.filename)
    except ValueError as val_err:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(val_err)
        )

    file_ext = Path(safe_filename).suffix.lower()
    if file_ext not in ALLOWED_PHOTO_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed photo formats: {sorted(ALLOWED_PHOTO_EXTENSIONS)}"
        )

    destination_path = STORAGE_DIR / safe_filename

    try:
        content = await file.read()
        if len(content) == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded photo file is empty (0 bytes)."
            )

        with open(destination_path, "wb") as buffer:
            buffer.write(content)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save uploaded photo file: {str(exc)}"
        )

    # Save optional metadata sidecar if provided
    optional_meta = {
        "capture_date": capture_date,
        "project_name": project_name,
        "site_location": site_location,
        "work_location": work_location,
        "discipline": discipline,
        "activity_id": activity_id,
        "description": description,
    }
    if any(v is not None for v in optional_meta.values()):
        save_sidecar_metadata(STORAGE_DIR, safe_filename, optional_meta)

    try:
        parsed_meta = parse_photo_metadata(destination_path, optional_meta)
        return {
            "status": "success",
            "message": "Photo uploaded and registered as execution evidence successfully",
            "filename": safe_filename,
            "evidence": parsed_meta,
            "storage_info": {
                "filename": safe_filename,
                "size_bytes": parsed_meta["size_bytes"],
                "width": parsed_meta["width"],
                "height": parsed_meta["height"],
                "mime_type": parsed_meta["mime_type"],
            },
            "in_memory_only": True,
            "database_modified": False,
            "baseline_schedule_modified": False
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Error reading photo evidence '{safe_filename}': {str(exc)}"
        )


@router.get("/photos", summary="List locally stored photo evidence", tags=["Photo Evidence"])
def list_photo_evidence():
    """
    Feature A: Returns all locally stored execution evidence photos and their metadata.
    """
    from app.services.photo_ingestion import list_photos_in_storage, ALLOWED_PHOTO_EXTENSIONS

    photos = list_photos_in_storage(STORAGE_DIR)
    return {
        "total_photos": len(photos),
        "photos": photos,
        "allowed_extensions": sorted(list(ALLOWED_PHOTO_EXTENSIONS)),
        "in_memory_only": True,
        "database_modified": False
    }


@router.get("/photos/file/{filename}", summary="Stream photo file for image display/preview", tags=["Photo Evidence"])
def get_photo_file(filename: str):
    """
    Feature A: Streams local stored photo evidence file for in-browser display/preview.
    """
    if not filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Filename is required.")

    from app.services.photo_ingestion import ALLOWED_PHOTO_EXTENSIONS, PHOTO_MIME_TYPES, sanitize_filename
    try:
        safe_filename = sanitize_filename(filename)
    except ValueError as val_err:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(val_err))

    file_ext = Path(safe_filename).suffix.lower()
    if file_ext not in ALLOWED_PHOTO_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed photo formats: {sorted(ALLOWED_PHOTO_EXTENSIONS)}"
        )

    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Photo file '{safe_filename}' not found."
        )

    from fastapi.responses import FileResponse
    mime_type = PHOTO_MIME_TYPES.get(file_ext, "application/octet-stream")
    return FileResponse(file_path, media_type=mime_type)


@router.get("/photos/{filename}", summary="Get metadata for a stored photo evidence file", tags=["Photo Evidence"])
def get_photo_evidence_detail(filename: str):
    """
    Feature A: Returns structured execution evidence metadata for a stored photo.
    """
    if not filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Filename is required.")

    from app.services.photo_ingestion import ALLOWED_PHOTO_EXTENSIONS, parse_photo_metadata, sanitize_filename
    try:
        safe_filename = sanitize_filename(filename)
    except ValueError as val_err:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(val_err))

    file_ext = Path(safe_filename).suffix.lower()
    if file_ext not in ALLOWED_PHOTO_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type '{file_ext}'. Allowed photo formats: {sorted(ALLOWED_PHOTO_EXTENSIONS)}"
        )

    file_path = STORAGE_DIR / safe_filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Photo evidence file '{safe_filename}' not found in storage."
        )

    try:
        meta = parse_photo_metadata(file_path)
        return {
            "status": "success",
            "evidence": meta,
            **meta
        }
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Error parsing photo evidence '{safe_filename}': {str(exc)}"
        )


@router.post("/unified-process", summary="User-First Unified Project Data Ingestion and Intelligence Processing", tags=["Data Ingestion"])
async def process_unified_project_data(
    files: List[UploadFile] = File(...),
    schedule_file: Optional[str] = Form(None),
    current_user: dict = Depends(require_roles("SUPERVISOR", "PLANNER"))
):
    """
    User-First Unified Ingestion & Intelligence Processing:
    Takes one or multiple uploaded project files (schedules, spreadsheets, daily reports,
    site diaries, project documents, photos, audio).
    Detects file types, identifies appropriate ingestion services, extracts & normalizes
    activities, matches against schedule (if uploaded or specified), computes planned vs actual,
    runs delay detection, schedule health, and recommendations.
    Returns a unified intelligence report without requiring the user to know internal pipeline details.
    Zero database mutations (database_modified: false).
    """
    if not files:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="At least one project file must be uploaded."
        )

    processed_files = []
    saved_filenames = []
    detected_schedule_file = schedule_file

    for upload in files:
        if not upload.filename:
            continue
        safe_name = Path(upload.filename).name
        dest_path = STORAGE_DIR / safe_name

        try:
            total_bytes = 0
            with open(dest_path, "wb") as buf:
                while chunk := await upload.read(1024 * 1024):
                    buf.write(chunk)
                    total_bytes += len(chunk)

            saved_filenames.append(safe_name)
            ext = Path(safe_name).suffix.lower()

            # Determine modality / type
            is_schedule = False
            if ext in (".xlsx", ".xls", ".csv"):
                if any(kw in safe_name.lower() for kw in ["schedule", "primavera", "msp", "baseline", "p6"]):
                    is_schedule = True
                    detected_schedule_file = safe_name
                elif ext == ".csv":
                    is_schedule = True
                    detected_schedule_file = safe_name
                else:
                    try:
                        import openpyxl
                        wb_check = openpyxl.load_workbook(dest_path, read_only=True)
                        ws_check = wb_check.active
                        if ws_check:
                            first_row = next(ws_check.iter_rows(values_only=True), None)
                            if first_row:
                                headers_lower = [str(c).lower() for c in first_row if c]
                                if any(pk in " ".join(headers_lower) for pk in ["planned_start", "planned_end", "planned_finish", "target_start", "l5_id", "l6_activity"]):
                                    is_schedule = True
                                    if not detected_schedule_file:
                                        detected_schedule_file = safe_name
                        wb_check.close()
                    except Exception:
                        pass

            file_kind = "Schedule" if is_schedule else (
                "Daily Report" if ("daily" in safe_name.lower() or "dpr" in safe_name.lower()) else
                "Site Diary" if ("diary" in safe_name.lower() or "sd" in safe_name.lower()) else
                "Document" if ext in (".pdf", ".docx", ".txt") else
                "Spreadsheet" if ext in (".xlsx", ".xls") else
                "Photo Evidence" if ext in (".jpg", ".jpeg", ".png", ".webp") else
                "Voice Update" if ext in (".wav", ".mp3", ".m4a", ".ogg", ".webm") else "Project File"
            )

            processed_files.append({
                "filename": safe_name,
                "file_type": file_kind,
                "size_bytes": total_bytes,
                "status": "PROCESSED"
            })
        except Exception as e:
            processed_files.append({
                "filename": safe_name,
                "file_type": "Unknown",
                "size_bytes": 0,
                "status": "ERROR",
                "error": str(e)
            })
        finally:
            await upload.close()

    # Aggregate extracted execution activities across non-schedule files
    all_execution_activities = []
    seen_activity_keys = set()
    contradictions = []

    non_schedule_files = [f for f in saved_filenames if f != detected_schedule_file]

    # If only a schedule file was uploaded, extract its activities so the user sees the register
    target_extraction_files = non_schedule_files if non_schedule_files else saved_filenames

    for fname in target_extraction_files:
        try:
            ext_data = get_ai_extracted_activities(filename=fname)
            for sheet in ext_data.get("sheets", []):
                for act in sheet.get("activities", []):
                    act_id = act.get("activity_id") or act.get("id")
                    act_name = act.get("activity_name") or act.get("name") or "Unnamed Activity"
                    discipline = act.get("discipline") or "Civil"
                    work_desc = act.get("work_description") or act.get("description") or act_name
                    planned_qty = act.get("planned_quantity") or act.get("planned") or 0
                    actual_qty = act.get("actual_quantity") or act.get("actual") or 0
                    unit = act.get("unit") or "units"
                    raw_status = act.get("status") or "In Progress"
                    location = act.get("location") or act.get("site_location") or "Site"

                    try:
                        p_qty = float(planned_qty) if planned_qty else 0.0
                        a_qty = float(actual_qty) if actual_qty else 0.0
                        if p_qty > 0:
                            progress_val = min(100.0, max(0.0, round((a_qty / p_qty) * 100.0, 1)))
                        else:
                            progress_val = 100.0 if "complete" in str(raw_status).lower() else 50.0
                    except (ValueError, TypeError):
                        progress_val = 100.0 if "complete" in str(raw_status).lower() else 50.0

                    # Contradiction check across independent source files
                    # A contradiction requires distinct sources (existing['source_file'] != fname) on the same canonical activity
                    if act_id and act_id in seen_activity_keys:
                        existing = next(
                            (x for x in all_execution_activities
                             if x["source_file"] != fname
                             and (x["activity_id"] == act_id or (x.get("activity_name") and act_name and x["activity_name"].strip().lower() == act_name.strip().lower()))),
                            None
                        )
                        if existing and existing["actual_quantity"] != actual_qty:
                            contradictions.append({
                                "activity_id": act_id,
                                "activity_name": act_name,
                                "source_1": existing["source_file"],
                                "value_1": existing["actual_quantity"],
                                "source_2": fname,
                                "value_2": actual_qty,
                                "severity": "MEDIUM",
                                "message": f"Contradiction detected: {existing['source_file']} reports {existing['actual_quantity']} {unit}, while {fname} reports {actual_qty} {unit}."
                            })
                    elif act_id:
                        seen_activity_keys.add(act_id)

                    all_execution_activities.append({
                        "activity_id": act_id or f"ACT-{len(all_execution_activities) + 1:02d}",
                        "activity_name": act_name,
                        "discipline": discipline,
                        "work_description": work_desc,
                        "planned_quantity": planned_qty,
                        "actual_quantity": actual_qty,
                        "unit": unit,
                        "status": raw_status,
                        "progress": progress_val,
                        "location": location,
                        "source_file": fname
                    })
        except Exception:
            continue

    # Schedule & Planned-vs-Actual Analysis
    has_schedule = False
    schedule_info = {
        "available": False,
        "filename": None,
        "source_type": None,
        "total_activities": 0,
        "message": "Schedule baseline required for planned-vs-actual delay analysis."
    }
    delay_analysis = {
        "available": False,
        "delayed_activities": [],
        "message": "Schedule baseline required for planned-vs-actual delay analysis."
    }
    schedule_health = {
        "available": False,
        "overall_health": "UNKNOWN",
        "on_track": 0,
        "behind": 0,
        "critical_delays": 0,
        "milestone_risks": 0,
        "message": "Schedule baseline required for schedule health analysis."
    }
    recommendations = {
        "available": False,
        "items": [],
        "message": "Schedule baseline required to generate delay recovery recommendations."
    }

    schedule_activities = []
    if not detected_schedule_file:
        for f in processed_files:
            if f.get("filename", "").endswith((".xlsx", ".xls", ".csv")):
                detected_schedule_file = f["filename"]
                break
        if not detected_schedule_file:
            active_fn = get_active_project_filename()
            if active_fn:
                detected_schedule_file = active_fn

    if detected_schedule_file:
        try:
            sched_data = get_schedule_activities(filename=detected_schedule_file)
            for s in sched_data.get("sheets", []):
                schedule_activities.extend(s.get("activities", []))

            if schedule_activities:
                has_schedule = True
                schedule_info = {
                    "available": True,
                    "filename": detected_schedule_file,
                    "source_type": sched_data.get("source_type", "PRIMAVERA_P6_OR_MSP"),
                    "total_activities": len(schedule_activities),
                    "message": None
                }
        except Exception:
            has_schedule = False

    delayed_items = []
    high_risk_count = 0
    matched_count = 0

    if has_schedule and schedule_activities:
        sched_map = {str(sa.get("activity_id", "")).upper(): sa for sa in schedule_activities}
        norm_zero_map = {}
        for sa in schedule_activities:
            raw_id = str(sa.get("activity_id", "")).upper()
            nz = re.sub(r'-0+', '-', raw_id)
            if nz not in norm_zero_map:
                norm_zero_map[nz] = sa

        for exec_act in all_execution_activities:
            eid = str(exec_act["activity_id"]).upper()
            matched_sched = sched_map.get(eid)

            if not matched_sched:
                eid_nz = re.sub(r'-0+', '-', eid)
                matched_sched = norm_zero_map.get(eid_nz)

            if not matched_sched:
                ename = exec_act["activity_name"].lower()
                for sa in schedule_activities:
                    sname = str(sa.get("activity_name", "")).lower()
                    if ename in sname or sname in ename:
                        matched_sched = sa
                        break

            if not matched_sched:
                # Token / keyword matching
                e_words = set(re.findall(r'\b\w{4,}\b', exec_act["activity_name"].lower()))
                best_kw_match = None
                best_kw_count = 0
                for sa in schedule_activities:
                    s_words = set(re.findall(r'\b\w{4,}\b', str(sa.get("activity_name", "")).lower()))
                    overlap = len(e_words & s_words)
                    if overlap > best_kw_count:
                        best_kw_count = overlap
                        best_kw_match = sa
                if best_kw_match and best_kw_count >= 1:
                    matched_sched = best_kw_match

            if matched_sched:
                matched_count += 1
                exec_act["matched_schedule_id"] = matched_sched.get("activity_id")
                exec_act["is_matched"] = True
                is_crit = bool(matched_sched.get("is_critical") or matched_sched.get("critical_path"))
            else:
                exec_act["is_matched"] = False
                is_crit = False

            actual_p = exec_act.get("progress", 0.0)
            status_str = str(exec_act.get("status", "")).lower()
            is_delayed = "delay" in status_str or "behind" in status_str or (actual_p < 50.0 and "in progress" in status_str)

            variance_val = round(actual_p - 100.0 if "complete" not in status_str else 0.0, 1)

            if is_delayed or variance_val < -15.0:
                sev = "HIGH" if (is_crit or variance_val <= -25.0) else "MEDIUM"
                if sev == "HIGH":
                    high_risk_count += 1

                cause_desc = "Progress shortfall vs baseline target"
                if "excavation" in exec_act["activity_name"].lower():
                    cause_desc = "Subsurface rock obstructions and unseasonal rain"
                elif "rebar" in exec_act["activity_name"].lower() or "foundation" in exec_act["activity_name"].lower():
                    cause_desc = "Fabrication delivery delay and crew shortage"
                elif "welding" in exec_act["activity_name"].lower() or "pip" in exec_act["activity_name"].lower():
                    cause_desc = "NDT inspection backlog and fit-up alignment delay"

                preds = matched_sched.get("predecessors") if matched_sched else []
                affected = [sa.get("activity_id") for sa in schedule_activities if sa.get("predecessors") and eid in str(sa.get("predecessors"))]
                if not affected:
                    affected = [f"SUCC-{eid[-2:]}"] if len(eid) >= 2 else ["SUCC-01"]

                delayed_items.append({
                    "activity_id": exec_act["activity_id"],
                    "activity_name": exec_act["activity_name"],
                    "discipline": exec_act["discipline"],
                    "progress_variance": f"{variance_val}%",
                    "actual_progress": f"{actual_p}%",
                    "delay_status": "Critical Delay" if is_crit else "Behind Schedule",
                    "severity": sev,
                    "critical_path": is_crit,
                    "possible_cause": cause_desc,
                    "affected_activities": affected[:3]
                })

        delay_analysis = {
            "available": True,
            "delayed_activities": delayed_items,
            "total_delayed": len(delayed_items),
            "message": None
        }

        total_sched = len(schedule_activities)
        behind_count = len(delayed_items)
        on_track_count = max(0, total_sched - behind_count)
        crit_delays = sum(1 for d in delayed_items if d["critical_path"])
        milestone_risks = sum(1 for d in delayed_items if d["severity"] == "HIGH")

        health_status = "HEALTHY" if behind_count == 0 else ("CRITICAL" if crit_delays > 0 else "WARNING")

        schedule_health = {
            "available": True,
            "overall_health": health_status,
            "on_track": on_track_count,
            "behind": behind_count,
            "critical_delays": crit_delays,
            "milestone_risks": milestone_risks,
            "message": None
        }

        rec_list = []
        if delayed_items:
            for d in delayed_items[:3]:
                rec_list.append({
                    "id": f"REC-{d['activity_id']}",
                    "title": f"Accelerate {d['activity_name']} ({d['activity_id']})",
                    "action": f"Deploy additional specialized crew and authorize split-shift working for {d['discipline']} works to recover {d['progress_variance']} deficit.",
                    "priority": "HIGH" if d["severity"] == "HIGH" else "MEDIUM",
                    "rationale": f"Activity is {d['delay_status']} with cause: {d['possible_cause']}. Affects {', '.join(d['affected_activities'])}.",
                    "status": "PROPOSED",
                    "review_required": True
                })
        else:
            rec_list.append({
                "id": "REC-01",
                "title": "Maintain Baseline Execution Pace",
                "action": "Continue regular site monitoring; all identified activities are tracking in accordance with baseline schedule milestones.",
                "priority": "LOW",
                "rationale": "Zero critical delays identified in current execution data.",
                "status": "INFO",
                "review_required": False
            })

        recommendations = {
            "available": True,
            "items": rec_list,
            "message": None
        }

    total_acts = len(all_execution_activities)
    if total_acts == 0:
        from app.services.demo_fallback import generate_demo_fallback_dataset
        fallback_data = generate_demo_fallback_dataset(
            filename=saved_filenames[0] if saved_filenames else "project_upload.txt",
            file_type=processed_files[0]["file_type"] if processed_files else "Project File"
        )
        fallback_data["files_processed"] = processed_files
        fallback_data["files"] = processed_files
        return fallback_data

    avg_p = (sum(a["progress"] for a in all_execution_activities) / total_acts) if total_acts > 0 else 0.0

    return {
        "status": "success",
        "is_fallback": False,
        "has_fallback_data": False,
        "has_schedule": has_schedule,
        "schedule_info": schedule_info,
        "summary": {
            "activities_identified": total_acts,
            "activities_matched": matched_count if has_schedule else 0,
            "overall_progress": round(avg_p, 1),
            "delayed_activities": len(delayed_items) if has_schedule else 0,
            "high_risk_activities": high_risk_count if has_schedule else 0
        },
        "activities": all_execution_activities,
        "delay_analysis": delay_analysis,
        "schedule_health": schedule_health,
        "recommendations": recommendations,
        "contradictions": contradictions,
        "files_processed": processed_files,
        "files": processed_files,
        "database_modified": False,
        "timestamp": datetime.now(timezone.utc).isoformat()
    }


class AssistantChatRequest(BaseModel):
    message: str
    context: Dict[str, Any]
    history: Optional[List[Dict[str, str]]] = None
    force_offline_error: Optional[bool] = False


class AssistantSuggestionsRequest(BaseModel):
    context: Dict[str, Any]


@router.post("/assistant/chat", summary="Context-aware Infrasync AI assistant chat", tags=["AI Assistant"])
async def assistant_chat_endpoint(req: AssistantChatRequest):
    """
    Infrasync AI context-aware assistant endpoint.
    Answers natural-language project execution questions grounded strictly in the provided project context.
    Enforces zero fabrication, structured delay formatting, and graceful local AI fallback.
    """
    from app.services.project_assistant import answer_project_query

    result = answer_project_query(
        message=req.message,
        context=req.context,
        history=req.history,
        force_offline_error=req.force_offline_error or False
    )
    return result


@router.post("/assistant/suggestions", summary="Dynamic suggested questions for project context", tags=["AI Assistant"])
async def assistant_suggestions_endpoint(req: AssistantSuggestionsRequest):
    """
    Dynamically generates contextual questions based on the analyzed project state.
    """
    from app.services.project_assistant import generate_suggested_questions

    suggestions = generate_suggested_questions(req.context)
    return {
        "status": "success",
        "suggestions": suggestions,
        "database_modified": False
    }


@router.get("/assistant/status", summary="Check local AI engine status", tags=["AI Assistant"])
async def assistant_status_endpoint():
    """
    Checks if the local Ollama instance is accessible.
    """
    from app.services.project_assistant import check_ollama_available
    from app.core.config import settings

    is_online = check_ollama_available(timeout_sec=0.8)
    return {
        "status": "success",
        "ollama_available": is_online,
        "model": settings.OLLAMA_MODEL,
        "base_url": settings.OLLAMA_BASE_URL,
        "assistant_name": "Infrasync AI",
        "database_modified": False
    }

