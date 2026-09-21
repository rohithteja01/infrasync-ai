import os
import io
import re
import csv
from pathlib import Path
from datetime import datetime, date
from typing import Dict, Any, List, Optional, Tuple, Union

import openpyxl

ALLOWED_SCHEDULE_EXTENSIONS = {".xlsx", ".xls", ".csv"}

# Comprehensive column candidates for deterministic mapping
COLUMN_ALIASES = {
    "activity_id": [
        "activity_id", "activity id", "activityid", "activity_code", "activity code",
        "act_id", "act id", "actid", "task_id", "task id", "taskid", "task_code",
        "code", "id", "activity_number", "act_num"
    ],
    "activity_name": [
        "activity_name", "activity name", "activityname", "task_name", "task name",
        "taskname", "activity_description", "activity description", "description",
        "task_desc", "task desc", "task", "activity", "name", "work_description",
        "scope_description"
    ],
    "wbs": [
        "wbs", "wbs_code", "wbs code", "wbscode", "wbs_name", "wbs name",
        "wbs_element", "wbs element", "work_breakdown_structure", "outline_number",
        "outline number", "wbs_number", "wbs_id"
    ],
    "level": [
        "level", "activity_level", "activity level", "wbs_level", "wbs level",
        "outline_level", "outline level", "hierarchical_level", "schedule_level", "lvl"
    ],
    "planned_start": [
        "planned_start", "planned start", "start_date", "start date", "target_start",
        "target start", "early_start", "early start", "baseline_start", "baseline start",
        "start", "commencement_date"
    ],
    "planned_finish": [
        "planned_finish", "planned finish", "finish_date", "finish date", "end_date",
        "end date", "target_finish", "target finish", "early_finish", "early finish",
        "baseline_finish", "baseline finish", "finish", "end", "completion_date"
    ],
    "duration": [
        "duration", "original_duration", "original duration", "planned_duration",
        "planned duration", "duration_days", "duration (days)", "dur", "days",
        "duration_d", "remaining_duration"
    ],
    "predecessors": [
        "predecessors", "predecessor", "predecessor_ids", "predecessor ids",
        "predecessor_id", "dependencies", "depends_on", "pred"
    ],
    "discipline": [
        "discipline", "trade", "department", "package", "engineering_discipline",
        "disc", "work_package", "resource_names"
    ],
    "progress_percent": [
        "%_complete", "% complete", "percent_complete", "percent complete",
        "physical_%_complete", "physical % complete", "%_work_complete",
        "% work complete", "progress", "progress_percent", "actual_%_complete"
    ],
    "status": [
        "activity_status", "activity status", "status", "state", "task_status"
    ]
}


def clean_text(value: Any) -> Optional[str]:
    """Cleans text values, removes redundant whitespace, and handles null indicators."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    val_str = str(value).strip()
    val_str = re.sub(r'\s+', ' ', val_str)
    if not val_str or val_str.lower() in ("nan", "none", "null", "n/a", "-", "--", "nil"):
        return None
    return val_str


def normalize_column_name(header: str) -> str:
    """Standardizes column header: strips punctuation, lowercase, snake_case."""
    if not header:
        return ""
    cleaned = str(header).strip()
    cleaned = re.sub(r'[\s\-\/\.\(\)]+', '_', cleaned)
    cleaned = re.sub(r'_+', '_', cleaned).strip('_')
    return cleaned.lower()


def parse_date_string(val: Any) -> Optional[str]:
    """Standardizes date value into ISO YYYY-MM-DD format."""
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null", "n/a", "-", "--"):
        return None

    # Try ISO YYYY-MM-DD first
    m_iso = re.match(r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})", s)
    if m_iso:
        y, m, d = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
        return f"{y:04d}-{m:02d}-{d:02d}"

    # Try common formats
    for fmt in (
        "%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y", "%d/%m/%y",
        "%d-%b-%Y", "%d-%b-%y", "%b %d, %Y", "%Y-%m-%d %H:%M:%S",
        "%d-%m-%Y", "%Y/%m/%d"
    ):
        try:
            date_part = s.split()[0] if " " in s and "-" in s else s
            dt = datetime.strptime(date_part, fmt)
            return dt.strftime("%Y-%m-%d")
        except Exception:
            pass

    return s


def parse_duration_numeric(val: Any) -> Optional[float]:
    """Extracts numeric duration in days (e.g. '41 days' -> 41.0, '13d' -> 13.0)."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null", "n/a", "-"):
        return None
    m = re.search(r"(\d+(?:\.\d+)?)", s)
    if m:
        try:
            num = float(m.group(1))
            if "hr" in s.lower() or "hour" in s.lower():
                num = round(num / 8.0, 2)
            return int(num) if num.is_integer() else num
        except Exception:
            pass
    return None


def format_duration(val: Any) -> Optional[str]:
    """Formats duration into human-readable string (e.g. '41 days')."""
    num = parse_duration_numeric(val)
    if num is None:
        raw_str = clean_text(val)
        return raw_str
    int_val = int(num) if isinstance(num, float) and num.is_integer() else num
    return f"{int_val} days" if int_val != 1 else "1 day"


def parse_progress_percent(val: Any) -> Optional[float]:
    """Parses progress percentage value (e.g. '70%', '0.70', 67 -> 70.0, 67.0)."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        v = float(val)
        if 0.0 < v <= 1.0:
            v = v * 100.0
        return round(v, 2)
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null", "n/a", "-"):
        return None
    m = re.search(r"(\d+(?:\.\d+)?)\s*%?", s)
    if m:
        try:
            v = float(m.group(1))
            if 0.0 < v <= 1.0 and "%" not in s:
                v = v * 100.0
            return round(v, 2)
        except Exception:
            pass
    return None


def normalize_schedule_level(raw_level: Any) -> str:
    """Classifies whether schedule activity is L5, L6, or another level."""
    cleaned = clean_text(raw_level)
    if not cleaned:
        return "Unspecified"

    cleaned_upper = cleaned.upper()
    if re.search(r'\bL[\s\-_]?5\b', cleaned_upper) or re.search(r'\bLEVEL[\s\-_]?5\b', cleaned_upper) or cleaned_upper == "5":
        return "L5"
    if re.search(r'\bL[\s\-_]?6\b', cleaned_upper) or re.search(r'\bLEVEL[\s\-_]?6\b', cleaned_upper) or cleaned_upper == "6":
        return "L6"
    match_other = re.search(r'\bL(?:EVEL)?[\s\-_]?([1-4])\b', cleaned_upper)
    if match_other:
        return f"L{match_other.group(1)}"
    if cleaned_upper in ("1", "2", "3", "4"):
        return f"L{cleaned_upper}"

    return cleaned


def detect_schedule_source_type(columns: List[str], sheet_name: str = "") -> str:
    """
    Deterministically identifies whether exported schedule is:
    - PRIMAVERA_P6_EXPORT
    - MS_PROJECT_EXPORT
    - GENERIC_SCHEDULE
    Based exclusively on column header patterns and sheet signatures.
    """
    normalized_headers = [re.sub(r'[\s\-_]+', ' ', col).strip().lower() for col in columns if col]
    combined_header_str = " ".join(normalized_headers)

    p6_score = 0
    msp_score = 0

    # Primavera P6 signatures
    if any(k in normalized_headers for k in ["activity id", "act id"]):
        p6_score += 2
    if "original duration" in normalized_headers or "original duration" in combined_header_str:
        p6_score += 2
    if "physical % complete" in normalized_headers or "physical % complete" in combined_header_str:
        p6_score += 2
    if "activity status" in normalized_headers or "activity status" in combined_header_str:
        p6_score += 2
    if any(k in normalized_headers for k in ["wbs code", "earned value", "planned value"]):
        p6_score += 1
    if any(k in sheet_name.upper() for k in ["TASK", "P6", "PRIMAVERA"]):
        p6_score += 2

    # MS Project signatures
    if any(k in normalized_headers for k in ["task name", "task id"]):
        msp_score += 2
    if any(k in normalized_headers for k in ["outline number", "outline level"]):
        msp_score += 3
    if "% work complete" in normalized_headers or "% work complete" in combined_header_str:
        msp_score += 2
    if any(k in normalized_headers for k in ["resource names", "predecessors"]):
        msp_score += 1
    if any(k in sheet_name.upper() for k in ["MS PROJECT", "MSP", "PROJECT"]):
        msp_score += 1

    if p6_score >= 3 and p6_score > msp_score:
        return "PRIMAVERA_P6_EXPORT"
    elif msp_score >= 3 and msp_score > p6_score:
        return "MS_PROJECT_EXPORT"
    else:
        return "GENERIC_SCHEDULE"


def detect_schedule_field_mapping(columns: List[str]) -> Dict[str, Optional[str]]:
    """
    Maps original source column headers to canonical schedule attributes.
    Returns: { canonical_field_name: original_source_column or None }
    """
    normalized_cols = {col: normalize_column_name(col) for col in columns if col}
    mapping: Dict[str, Optional[str]] = {canonical: None for canonical in COLUMN_ALIASES}
    used_columns = set()

    # Exact matches
    for canonical, aliases in COLUMN_ALIASES.items():
        for original_col, norm_col in normalized_cols.items():
            if original_col in used_columns:
                continue
            cleaned_norm = norm_col.replace("_", " ")
            for alias in aliases:
                alias_norm = alias.replace("_", " ").lower()
                if cleaned_norm == alias_norm or norm_col == alias.replace(" ", "_").lower():
                    mapping[canonical] = original_col
                    used_columns.add(original_col)
                    break
            if mapping[canonical] is not None:
                break

    # Substring / partial matches
    for canonical, aliases in COLUMN_ALIASES.items():
        if mapping[canonical] is not None:
            continue
        for original_col, norm_col in normalized_cols.items():
            if original_col in used_columns:
                continue
            cleaned_norm = norm_col.replace("_", " ")
            for alias in aliases:
                alias_norm = alias.replace("_", " ").lower()
                if (alias_norm in cleaned_norm or cleaned_norm in alias_norm) and len(alias_norm) > 2:
                    mapping[canonical] = original_col
                    used_columns.add(original_col)
                    break
            if mapping[canonical] is not None:
                break

    return mapping


def detect_header_row(rows: List[List[Any]]) -> Tuple[int, List[str]]:
    """
    Scans the top rows to find the most probable header row containing schedule attributes.
    Returns (header_row_index, list_of_column_names).
    """
    best_index = 0
    best_score = 0
    best_headers: List[str] = []

    for idx, row in enumerate(rows[:15]):
        if not row:
            continue
        row_str_cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if len(row_str_cells) < 2:
            continue

        score = 0
        for cell in row_str_cells:
            cell_norm = normalize_column_name(cell)
            for aliases in COLUMN_ALIASES.values():
                if any(cell_norm == a.replace(" ", "_") or cell_norm == a for a in aliases):
                    score += 1
                    break

        if score > best_score:
            best_score = score
            best_index = idx
            best_headers = [str(c).strip() if c is not None else f"Column_{i+1}" for i, c in enumerate(row)]

    if best_score == 0 and rows:
        best_headers = [str(c).strip() if c is not None else f"Column_{i+1}" for i, c in enumerate(rows[0])]
        best_index = 0

    return best_index, best_headers


def parse_schedule_rows(
    sheet_name: str,
    columns: List[str],
    raw_records: List[Dict[str, Any]],
    filename: str,
    source_format: str
) -> Dict[str, Any]:
    """
    Parses row records from a worksheet or CSV into standardized schedule format.
    Preserves all unmapped source columns under 'original_attributes'.
    """
    field_mapping = detect_schedule_field_mapping(columns)

    id_col = field_mapping.get("activity_id")
    name_col = field_mapping.get("activity_name")
    wbs_col = field_mapping.get("wbs")
    level_col = field_mapping.get("level")
    start_col = field_mapping.get("planned_start")
    finish_col = field_mapping.get("planned_finish")
    dur_col = field_mapping.get("duration")
    pred_col = field_mapping.get("predecessors")
    disc_col = field_mapping.get("discipline")
    prog_col = field_mapping.get("progress_percent")
    stat_col = field_mapping.get("status")

    mapped_cols_set = {v for v in field_mapping.values() if v is not None}

    activities: List[Dict[str, Any]] = []
    l5_count = 0
    l6_count = 0
    other_count = 0

    for idx, row in enumerate(raw_records):
        raw_dict = row or {}
        # Skip completely empty rows
        if not any(v is not None and str(v).strip() for v in raw_dict.values()):
            continue

        raw_id = clean_text(raw_dict.get(id_col)) if id_col else None
        raw_name = clean_text(raw_dict.get(name_col)) if name_col else None

        # At least activity ID or activity name must exist for a valid row
        if not raw_id and not raw_name:
            continue

        act_id = raw_id or f"SCHED-{idx + 1:03d}"
        act_name = raw_name or f"Activity {idx + 1}"

        wbs = clean_text(raw_dict.get(wbs_col)) if wbs_col else None
        raw_lvl = clean_text(raw_dict.get(level_col)) if level_col else None
        level = normalize_schedule_level(raw_lvl)

        p_start = parse_date_string(raw_dict.get(start_col)) if start_col else None
        p_finish = parse_date_string(raw_dict.get(finish_col)) if finish_col else None

        raw_dur = raw_dict.get(dur_col) if dur_col else None
        duration_num = parse_duration_numeric(raw_dur)
        duration_str = format_duration(raw_dur) if raw_dur is not None else None

        predecessors = clean_text(raw_dict.get(pred_col)) if pred_col else None
        discipline = clean_text(raw_dict.get(disc_col)) if disc_col else "General"

        prog_pct = parse_progress_percent(raw_dict.get(prog_col)) if prog_col else None
        status = clean_text(raw_dict.get(stat_col)) if stat_col else None

        if not status:
            if prog_pct is not None:
                if prog_pct >= 100.0:
                    status = "Completed"
                elif prog_pct > 0:
                    status = "In Progress"
                else:
                    status = "Not Started"

        # Level tallies
        if level == "L5":
            l5_count += 1
        elif level == "L6":
            l6_count += 1
        else:
            other_count += 1

        # Preserve unmapped columns under original_attributes
        original_attributes = {
            k: clean_text(v)
            for k, v in raw_dict.items()
            if k not in mapped_cols_set and clean_text(v) is not None
        }

        activity_record = {
            "index": idx + 1,
            "activity_id": act_id,
            "activity_name": act_name,
            "wbs": wbs or "-",
            "level": level,
            "raw_level": raw_lvl or "-",
            "planned_start": p_start or "-",
            "planned_finish": p_finish or "-",
            "duration": duration_str or "-",
            "duration_days": duration_num,
            "discipline": discipline or "General",
            "predecessors": predecessors or "-",
            "progress_percent": prog_pct,
            "status": status,
            "original_attributes": original_attributes,
            "source_file": filename,
            "source_format": source_format,
            "raw_values": {k: clean_text(v) for k, v in raw_dict.items()}
        }
        activities.append(activity_record)

    return {
        "sheet_name": sheet_name,
        "total_schedule_activities": len(activities),
        "l5_count": l5_count,
        "l6_count": l6_count,
        "other_count": other_count,
        "field_mapping": field_mapping,
        "activities": activities
    }


def parse_schedule_file(file_input: Union[str, Path, bytes], filename: Optional[str] = None) -> Dict[str, Any]:
    """
    Parses any supported schedule export (.xlsx, .xls, .csv).
    Accepts either a file path (str or Path) or raw binary content (bytes).
    Detects source type (PRIMAVERA_P6_EXPORT, MS_PROJECT_EXPORT, GENERIC_SCHEDULE),
    normalizes schedule activities, and returns standard representation in-memory.
    Never modifies database or schedule baseline.
    """
    if isinstance(file_input, (bytes, bytearray)):
        if len(file_input) == 0:
            raise ValueError(f"Schedule file '{filename or 'upload'}' is empty (0 bytes).")
        safe_filename = Path(filename).name if filename else "schedule_export.xlsx"
        ext = Path(safe_filename).suffix.lower()
        if ext not in ALLOWED_SCHEDULE_EXTENSIONS:
            raise ValueError(f"Unsupported schedule format '{ext}'. Allowed: {sorted(ALLOWED_SCHEDULE_EXTENSIONS)}")
        file_bytes = bytes(file_input)
        is_bytes = True
        stem_name = Path(safe_filename).stem
    else:
        p = Path(file_input).resolve()
        if not p.exists() or not p.is_file():
            raise FileNotFoundError(f"Schedule file '{p.name}' not found.")
        if p.stat().st_size == 0:
            raise ValueError(f"Schedule file '{p.name}' is empty (0 bytes).")
        ext = p.suffix.lower()
        if ext not in ALLOWED_SCHEDULE_EXTENSIONS:
            raise ValueError(f"Unsupported schedule format '{ext}'. Allowed: {sorted(ALLOWED_SCHEDULE_EXTENSIONS)}")
        safe_filename = p.name
        stem_name = p.stem
        is_bytes = False
        file_bytes = None

    source_format = ext.lstrip(".")
    sheets_output: List[Dict[str, Any]] = []
    all_activities: List[Dict[str, Any]] = []
    detected_source_type = "GENERIC_SCHEDULE"

    if ext == ".csv":
        # Read CSV file with encoding fallback
        raw_text = None
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                if is_bytes:
                    raw_text = file_bytes.decode(enc)
                else:
                    with open(p, "r", encoding=enc, errors="replace") as f:
                        raw_text = f.read()
                break
            except Exception:
                continue

        if not raw_text or not raw_text.strip():
            raise ValueError(f"CSV file '{safe_filename}' contains no readable data.")

        # Detect delimiter
        first_line = raw_text.splitlines()[0] if raw_text.splitlines() else ""
        delim = ","
        if ";" in first_line and first_line.count(";") > first_line.count(","):
            delim = ";"
        elif "\t" in first_line and first_line.count("\t") > first_line.count(","):
            delim = "\t"

        reader = csv.reader(raw_text.splitlines(), delimiter=delim)
        all_rows = list(reader)
        if not all_rows:
            raise ValueError(f"No rows found in CSV schedule '{safe_filename}'.")

        h_idx, headers = detect_header_row(all_rows)
        data_rows = all_rows[h_idx + 1:]

        records: List[Dict[str, Any]] = []
        for r in data_rows:
            rec = {}
            for col_i, h in enumerate(headers):
                rec[h] = r[col_i] if col_i < len(r) else None
            records.append(rec)

        detected_source_type = detect_schedule_source_type(headers, safe_filename)
        parsed_sheet = parse_schedule_rows(
            sheet_name=stem_name,
            columns=headers,
            raw_records=records,
            filename=safe_filename,
            source_format=source_format
        )
        sheets_output.append(parsed_sheet)
        all_activities.extend(parsed_sheet["activities"])

    elif ext == ".xlsx":
        wb_source = io.BytesIO(file_bytes) if is_bytes else p
        wb = openpyxl.load_workbook(wb_source, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            raw_rows = list(ws.iter_rows(values_only=True))
            if not raw_rows:
                continue

            h_idx, headers = detect_header_row(raw_rows)
            data_rows = raw_rows[h_idx + 1:]

            records = []
            for r in data_rows:
                rec = {}
                for col_i, h in enumerate(headers):
                    rec[h] = r[col_i] if col_i < len(r) else None
                records.append(rec)

            st = detect_schedule_source_type(headers, sname)
            if st != "GENERIC_SCHEDULE" or detected_source_type == "GENERIC_SCHEDULE":
                detected_source_type = st

            parsed_sheet = parse_schedule_rows(
                sheet_name=sname,
                columns=headers,
                raw_records=records,
                filename=safe_filename,
                source_format=source_format
            )
            sheets_output.append(parsed_sheet)
            all_activities.extend(parsed_sheet["activities"])
        wb.close()

    elif ext == ".xls":
        try:
            import xlrd
            if is_bytes:
                wb = xlrd.open_workbook(file_contents=file_bytes)
            else:
                wb = xlrd.open_workbook(p)
            for sname in wb.sheet_names():
                sh = wb.sheet_by_name(sname)
                raw_rows = []
                for r_i in range(sh.nrows):
                    raw_rows.append(sh.row_values(r_i))

                if not raw_rows:
                    continue

                h_idx, headers = detect_header_row(raw_rows)
                data_rows = raw_rows[h_idx + 1:]

                records = []
                for r in data_rows:
                    rec = {}
                    for col_i, h in enumerate(headers):
                        rec[h] = r[col_i] if col_i < len(r) else None
                    records.append(rec)

                st = detect_schedule_source_type(headers, sname)
                if st != "GENERIC_SCHEDULE" or detected_source_type == "GENERIC_SCHEDULE":
                    detected_source_type = st

                parsed_sheet = parse_schedule_rows(
                    sheet_name=sname,
                    columns=headers,
                    raw_records=records,
                    filename=safe_filename,
                    source_format=source_format
                )
                sheets_output.append(parsed_sheet)
                all_activities.extend(parsed_sheet["activities"])
        except ImportError:
            raise RuntimeError(f"xlrd library not available to parse legacy XLS file '{safe_filename}'.")

    if not all_activities:
        raise ValueError(f"No usable schedule activities could be extracted from '{safe_filename}'.")

    total_acts = sum(s["total_schedule_activities"] for s in sheets_output)
    total_l5 = sum(s["l5_count"] for s in sheets_output)
    total_l6 = sum(s["l6_count"] for s in sheets_output)
    total_other = sum(s["other_count"] for s in sheets_output)

    return {
        "success": True,
        "status": "success",
        "filename": safe_filename,
        "source_type": detected_source_type,
        "source_format": source_format,
        "total_activities": total_acts,
        "total_schedule_activities": total_acts,
        "sheet_count": len(sheets_output),
        "l5_count": total_l5,
        "l6_count": total_l6,
        "other_count": total_other,
        "sheets": sheets_output,
        "activities": all_activities,
        "in_memory_only": True,
        "database_modified": False,
        "baseline_schedule_modified": False
    }


def generate_sample_primavera_export(target_path: Path) -> Path:
    """
    Generates a realistic Primavera P6 export (.xlsx) containing L5/L6 activities
    with typical Primavera column naming, activity codes, and attributes.
    """
    target_path = Path(target_path).resolve()
    target_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TASK"

    headers = [
        "Activity ID", "Activity Name", "WBS Code", "Activity Level",
        "Planned Start", "Planned Finish", "Original Duration",
        "Discipline", "Predecessors", "Physical % Complete",
        "Activity Status", "Calendar", "Total Float (d)"
    ]
    ws.append(headers)

    p6_rows = [
        ("CIV-L5-01", "Main Substructure Construction", "1.1.2.1", "L5", "2025-01-05", "2025-02-15", "41d", "Civil", None, "0%", "Not Started", "Standard 6-Day", 0),
        ("CIV-L6-01", "Subgrade excavation and compaction", "1.1.2.1.1", "L6", "2025-01-05", "2025-01-18", "13d", "Civil", None, "70%", "In Progress", "Standard 6-Day", 0),
        ("CIV-L6-02", "Foundation rebar fixing and shuttering", "1.1.2.1.2", "L6", "2025-01-19", "2025-01-30", "11d", "Civil", "CIV-L6-01", "67%", "In Progress", "Standard 6-Day", 0),
        ("CIV-L6-03", "M25 grade raft concrete pouring", "1.1.2.1.3", "L6", "2025-02-01", "2025-02-15", "14d", "Civil", "CIV-L6-02", "0%", "Not Started", "Standard 6-Day", 0),
        ("PIP-L5-01", "Process Header Piping Fabrication", "1.2.1.0", "L5", "2025-02-10", "2025-03-25", "43d", "Piping", None, "0%", "Not Started", "Standard 6-Day", 5),
        ("PIP-L6-01", "12-inch carbon steel pipe spool pre-fabrication", "1.2.1.0.1", "L6", "2025-02-10", "2025-02-28", "18d", "Piping", None, "100%", "Completed", "Standard 6-Day", 5),
        ("PIP-L6-02", "Pipe spool fit-up and butt welding", "1.2.1.0.2", "L6", "2025-03-01", "2025-03-15", "14d", "Piping", "PIP-L6-01", "0%", "Not Started", "Standard 6-Day", 5),
        ("PIP-L6-03", "Hydrostatic pressure testing (25 bar)", "1.2.1.0.3", "L6", "2025-03-16", "2025-03-25", "9d", "Piping", "PIP-L6-02", "0%", "Not Started", "Standard 6-Day", 5),
        ("MEC-L5-01", "Compressor Package Mechanical Installation", "1.3.4.0", "L5", "2025-03-01", "2025-04-10", "40d", "Mechanical", None, "0%", "Not Started", "Standard 6-Day", 0),
        ("MEC-L6-01", "Base plate grouting and alignment", "1.3.4.0.1", "L6", "2025-03-01", "2025-03-12", "11d", "Mechanical", None, "0%", "Not Started", "Standard 6-Day", 0),
        ("MEC-L6-02", "Centrifugal compressor skid heavy rigging", "1.3.4.0.2", "L6", "2025-03-13", "2025-03-28", "15d", "Mechanical", "MEC-L6-01", "0%", "Not Started", "Standard 6-Day", 0),
        ("MEC-L6-03", "Driver-compressor shaft laser alignment", "1.3.4.0.3", "L6", "2025-03-29", "2025-04-10", "12d", "Mechanical", "MEC-L6-02", "0%", "Not Started", "Standard 6-Day", 0),
    ]

    for row in p6_rows:
        ws.append(list(row))

    wb.save(target_path)
    wb.close()
    return target_path


def generate_sample_ms_project_export(target_path: Path) -> Path:
    """
    Generates a realistic MS Project export (.xlsx) containing L5/L6 activities
    with typical Microsoft Project column naming and structure.
    """
    target_path = Path(target_path).resolve()
    target_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task_Sheet"

    headers = [
        "Task ID", "Task Name", "Outline Number", "Outline Level",
        "Start Date", "Finish Date", "Duration", "Predecessors",
        "% Work Complete", "Resource Names", "Milestone", "Critical"
    ]
    ws.append(headers)

    msp_rows = [
        ("CIV-L5-01", "Main Substructure Construction", "1.1.2.1", "Level 5", "2025-01-05", "2025-02-15", "41 days", None, "0%", "Civil Crew Lead", "No", "Yes"),
        ("CIV-L6-01", "Subgrade excavation and compaction", "1.1.2.1.1", "Level 6", "2025-01-05", "2025-01-18", "13 days", None, "70%", "Civil Crew A", "No", "Yes"),
        ("CIV-L6-02", "Foundation rebar fixing and shuttering", "1.1.2.1.2", "Level 6", "2025-01-19", "2025-01-30", "11 days", "CIV-L6-01", "67%", "Civil Crew B", "No", "Yes"),
        ("CIV-L6-03", "M25 grade raft concrete pouring", "1.1.2.1.3", "Level 6", "2025-02-01", "2025-02-15", "14 days", "CIV-L6-02", "0%", "Concrete Gang", "No", "Yes"),
        ("PIP-L5-01", "Process Header Piping Fabrication", "1.2.1.0", "Level 5", "2025-02-10", "2025-03-25", "43 days", None, "0%", "Piping Superintendent", "No", "No"),
        ("PIP-L6-01", "12-inch carbon steel pipe spool pre-fabrication", "1.2.1.0.1", "Level 6", "2025-02-10", "2025-02-28", "18 days", None, "100%", "Welders Team 1", "No", "No"),
        ("PIP-L6-02", "Pipe spool fit-up and butt welding", "1.2.1.0.2", "Level 6", "2025-03-01", "2025-03-15", "14 days", "PIP-L6-01", "0%", "Fit-up Crew", "No", "No"),
        ("PIP-L6-03", "Hydrostatic pressure testing (25 bar)", "1.2.1.0.3", "Level 6", "2025-03-16", "2025-03-25", "9 days", "PIP-L6-02", "0%", "Testing Crew", "No", "No"),
        ("MEC-L5-01", "Compressor Package Mechanical Installation", "1.3.4.0", "Level 5", "2025-03-01", "2025-04-10", "40 days", None, "0%", "Rigging Lead", "No", "Yes"),
        ("MEC-L6-01", "Base plate grouting and alignment", "1.3.4.0.1", "Level 6", "2025-03-01", "2025-03-12", "11 days", None, "0%", "Millwright Team", "No", "Yes"),
        ("MEC-L6-02", "Centrifugal compressor skid heavy rigging", "1.3.4.0.2", "Level 6", "2025-03-13", "2025-03-28", "15 days", "MEC-L6-01", "0%", "Crane Team", "No", "Yes"),
        ("MEC-L6-03", "Driver-compressor shaft laser alignment", "1.3.4.0.3", "Level 6", "2025-03-29", "2025-04-10", "12 days", "MEC-L6-02", "0%", "Specialist Team", "No", "Yes"),
    ]

    for row in msp_rows:
        ws.append(list(row))

    wb.save(target_path)
    wb.close()
    return target_path


def generate_sample_csv_export(target_path: Path) -> Path:
    """
    Generates a realistic standard CSV schedule export.
    """
    target_path = Path(target_path).resolve()
    target_path.parent.mkdir(parents=True, exist_ok=True)

    headers = [
        "Activity ID", "Activity Name", "WBS", "Level",
        "Planned Start", "Planned Finish", "Duration",
        "Discipline", "Predecessors", "% Complete"
    ]

    rows = [
        ["CIV-L5-01", "Main Substructure Construction", "1.1.2.1", "L5", "2025-01-05", "2025-02-15", "41", "Civil", "", "0%"],
        ["CIV-L6-01", "Subgrade excavation and compaction", "1.1.2.1.1", "L6", "2025-01-05", "2025-01-18", "13", "Civil", "", "70%"],
        ["CIV-L6-02", "Foundation rebar fixing and shuttering", "1.1.2.1.2", "L6", "2025-01-19", "2025-01-30", "11", "Civil", "CIV-L6-01", "67%"],
        ["CIV-L6-03", "M25 grade raft concrete pouring", "1.1.2.1.3", "L6", "2025-02-01", "2025-02-15", "14", "Civil", "CIV-L6-02", "0%"],
        ["PIP-L5-01", "Process Header Piping Fabrication", "1.2.1.0", "L5", "2025-02-10", "2025-03-25", "43", "Piping", "", "0%"],
        ["PIP-L6-01", "12-inch carbon steel pipe spool pre-fabrication", "1.2.1.0.1", "L6", "2025-02-10", "2025-02-28", "18", "Piping", "", "100%"],
        ["PIP-L6-02", "Pipe spool fit-up and butt welding", "1.2.1.0.2", "L6", "2025-03-01", "2025-03-15", "14", "Piping", "PIP-L6-01", "0%"],
        ["PIP-L6-03", "Hydrostatic pressure testing (25 bar)", "1.2.1.0.3", "L6", "2025-03-16", "2025-03-25", "9", "Piping", "PIP-L6-02", "0%"],
        ["MEC-L5-01", "Compressor Package Mechanical Installation", "1.3.4.0", "L5", "2025-03-01", "2025-04-10", "40", "Mechanical", "", "0%"],
        ["MEC-L6-01", "Base plate grouting and alignment", "1.3.4.0.1", "L6", "2025-03-01", "2025-03-12", "11", "Mechanical", "", "0%"],
        ["MEC-L6-02", "Centrifugal compressor skid heavy rigging", "1.3.4.0.2", "L6", "2025-03-13", "2025-03-28", "15", "Mechanical", "MEC-L6-01", "0%"],
        ["MEC-L6-03", "Driver-compressor shaft laser alignment", "1.3.4.0.3", "L6", "2025-03-29", "2025-04-10", "12", "Mechanical", "MEC-L6-02", "0%"],
    ]

    with open(target_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)

    return target_path


# Function aliases for compatibility
detect_source_format = detect_schedule_source_type
map_schedule_columns = detect_schedule_field_mapping

